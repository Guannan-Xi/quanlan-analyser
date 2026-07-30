"""Build a customer-facing delivery from the real ernie discrete TI search."""

from __future__ import annotations

import csv
import hashlib
import html
import json
import shutil
import subprocess
import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from simnibs_inverse_report_content import REFERENCES, render_customer_report


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "simnibs_inverse_ti_discrete_ernie_20260729"
RAW = OUT / "raw"
FIG = OUT / "figures"
TABLE = OUT / "tables"
METHOD = OUT / "methods"
REPRO = OUT / "reproduction"
M2M = ROOT / "data" / "simnibs_examples_v4_1" / "extracted" / "m2m_ernie"
SIMNIBS_PYTHON = Path(r"C:\Users\Administrator\Miniconda3\envs\simnibs_env\python.exe")
FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")
CONDUCTIVITIES_S_PER_M = {
    "WM": 0.126, "GM": 0.275, "CSF": 1.654, "Bone": 0.010,
    "Scalp": 0.465, "Eye_balls": 0.500, "Compact_bone": 0.008,
    "Spongy_bone": 0.025, "Blood": 0.600, "Muscle": 0.160,
    "Cartilage": 0.880, "Fat": 0.078,
}
ELECTRODE_MATERIAL_MODEL = {
    "simnibs_definition": "simple_electrode_single_layer",
    "layer_count": 1,
    "layers": [{
        "material": "Electrode_rubber",
        "thickness_mm": 2.0,
        "conductivity_s_per_m": 29.4,
    }],
    "gel_layer_modeled": False,
    "sponge_layer_modeled": False,
    "contact_impedance_modeled": False,
    "evidence": [
        "reproduction/scripts/run_simnibs_inverse_ti_discrete_ernie.py",
        "reproduction/scripts/validate_simnibs_inverse_ti_winner.py",
        "raw/fem/simnibs_simulation_20260729-024122.log",
        "SimNIBS 4.6.0 default TDCSLIST conductivity: Electrode_rubber=29.4 S/m",
    ],
    "transfer_boundary": "A customer electrode with gel, sponge, contact impedance, or different material properties requires a new layered electrode model and FEM solve.",
}
SEARCH_MESH = {"nodes": 801_484, "tetrahedra": 4_485_786, "triangles": 1_129_073}
INDEPENDENT_MESH = {"nodes": 798_697, "tetrahedra": 4_470_660, "triangles": 1_118_989}
ELECTRODE_GEOMETRY = [
    ("F5", "FC3", 40.209757059629744, 0.20975705962974445),
    ("P5", "CP3", 46.84902541334107, 6.849025413341067),
    ("F5", "P5", 126.5149092589389, 86.5149092589389),
    ("FC3", "CP3", 71.40000381907925, 31.400003819079245),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def font(size: int, bold: bool = False):
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT), size)


def save_chart(filename: str, title: str, subtitle: str, labels: list[str], values: list[float], unit: str, color: str) -> None:
    width, height = 2000, 1200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((120, 80), title, fill="#17211D", font=font(54, True))
    draw.text((120, 155), subtitle, fill="#5E6B64", font=font(28))
    left, top, right, bottom = 260, 300, 1840, 1010
    draw.line((left, bottom, right, bottom), fill="#89958F", width=3)
    maximum = max(values) * 1.15 if values else 1
    gap = (right - left) / len(values)
    for index, (label, value) in enumerate(zip(labels, values)):
        x0 = left + index * gap + gap * 0.18
        x1 = left + (index + 1) * gap - gap * 0.18
        y0 = bottom - (bottom - top) * value / maximum
        draw.rectangle((x0, y0, x1, bottom), fill=color)
        draw.text((x0, y0 - 54), f"{value:.4f}", fill="#17211D", font=font(28, True))
        box = draw.textbbox((0, 0), label, font=font(25))
        draw.text(((x0 + x1 - (box[2] - box[0])) / 2, bottom + 28), label, fill="#17211D", font=font(25))
    draw.text((120, 1095), f"单位：{unit}。柱高来自 report_data.json 的未四舍五入数值。", fill="#5E6B64", font=font(24))
    image.save(FIG / f"{filename}.png", dpi=(300, 300))
    bars = []
    for index, (label, value) in enumerate(zip(labels, values)):
        x0 = 260 + index * ((1840 - 260) / len(values)) + 50
        h = 620 * value / maximum
        bars.append(f'<rect x="{x0:.1f}" y="{1010-h:.1f}" width="{(1840-260)/len(values)-100:.1f}" height="{h:.1f}" fill="{color}"/><text x="{x0:.1f}" y="{1010-h-14:.1f}" font-size="26">{value:.4f}</text><text x="{x0:.1f}" y="1060" font-size="24">{html.escape(label)}</text>')
    bar_markup = "".join(bars)
    svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="2000" height="1200" viewBox="0 0 2000 1200"><rect width="2000" height="1200" fill="white"/><text x="120" y="125" font-size="54" font-weight="700">{html.escape(title)}</text><text x="120" y="185" font-size="28" fill="#5E6B64">{html.escape(subtitle)}</text><line x1="260" y1="1010" x2="1840" y2="1010" stroke="#89958F" stroke-width="3"/>{bar_markup}<text x="120" y="1140" font-size="24" fill="#5E6B64">单位：{html.escape(unit)}</text></svg>'
    (FIG / f"{filename}.svg").write_text(svg, encoding="utf-8")


def workbook(result: dict, validation: dict, figure_readiness: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active; ws.title = "Readme"
    ws.append(["字段", "内容"])
    for row in (("分析模式", "受限候选库内的离散 TI 逆向搜索"), ("数据", result["data_source"]), ("边界", result["execution_boundary"]), ("临床边界", result["clinical_boundary"])):
        ws.append(row)
    ws.append(["最终结果工作表", "Final four-electrode metrics；用于论文引用和二次统计"])
    ws.append(["搜索阶段工作表", "Search-stage metrics；仅用于候选排序，不得与最终结果跨表拼接"])
    ws.append(["阶段对照工作表", "Search-final comparison；同时报告 ROI mean、靶外 mean 和目标/靶外均值比的变化"])
    ranking = wb.create_sheet("Candidate ranking")
    columns = list(result["all_candidates"][0])
    ranking.append(columns)
    for row in result["all_candidates"]:
        ranking.append([row.get(key) for key in columns])
    metrics = wb.create_sheet("Search-stage metrics")
    metric_keys = list(result["metrics"]["target"])
    metrics.append(["region", *metric_keys])
    for region, values in result["metrics"].items():
        metrics.append([region, *[values[key] for key in metric_keys]])
    independent = wb.create_sheet("Final four-electrode metrics")
    independent.append(["region", *metric_keys])
    for region, values in validation["independent_metrics"].items():
        independent.append([region, *[values[key] for key in metric_keys]])
    independent.append([])
    independent.append(["metric", "value"])
    independent.append(["target_to_off_target_mean_ratio", validation["independent_target_to_off_target_mean_ratio"]])
    independent.append(["comparison_threshold_v_per_m", validation["comparison_threshold_v_per_m"]])
    sensitivity = validation["recompute_half_max_threshold_sensitivity"]
    independent.append(["recompute_half_max_threshold_v_per_m", sensitivity["threshold_v_per_m"]])
    independent.append(["sensitivity_target_suprathreshold_volume_mm3", sensitivity["target"]["suprathreshold_volume_mm3"]])
    independent.append(["sensitivity_target_threshold_coverage_pct", sensitivity["target"]["threshold_coverage_pct"]])
    independent.append(["sensitivity_off_target_suprathreshold_volume_mm3", sensitivity["off_target"]["suprathreshold_volume_mm3"]])
    independent.append(["sensitivity_off_target_threshold_coverage_pct", sensitivity["off_target"]["threshold_coverage_pct"]])
    independent.append(["sensitivity_target_share_of_all_suprathreshold_gray_pct", sensitivity["target_share_of_all_suprathreshold_gray_pct"]])
    comparison = wb.create_sheet("Search-final comparison")
    comparison.append(["metric", "search_stage", "four_electrode_recompute", "relative_change_pct", "interpretation"])
    search_roi_mean = validation["search_stage_roi_mean_v_per_m"]
    final_roi_mean = validation["independent_rerun_roi_mean_v_per_m"]
    search_off_target_mean = result["metrics"]["off_target"]["mean"]
    final_off_target_mean = validation["independent_metrics"]["off_target"]["mean"]
    search_mean_ratio = result["winner"]["target_to_off_target_mean_ratio"]
    final_mean_ratio = validation["independent_target_to_off_target_mean_ratio"]
    comparison_boundary = "descriptive cross-stage difference; configuration attribution not identifiable because the search-stage calibration-error summary is unavailable"
    comparison.append(["target_roi_mean_v_per_m", search_roi_mean, final_roi_mean, 100.0 * (final_roi_mean - search_roi_mean) / search_roi_mean, comparison_boundary])
    comparison.append(["off_target_mean_v_per_m", search_off_target_mean, final_off_target_mean, 100.0 * (final_off_target_mean - search_off_target_mean) / search_off_target_mean, comparison_boundary])
    comparison.append(["target_to_off_target_mean_ratio", search_mean_ratio, final_mean_ratio, 100.0 * (final_mean_ratio - search_mean_ratio) / search_mean_ratio, comparison_boundary])
    constraints = wb.create_sheet("Constraints")
    constraints.append(["constraint", "value"])
    for key, value in result["constraints"].items(): constraints.append([key, json.dumps(value, ensure_ascii=False)])
    provenance = wb.create_sheet("Provenance")
    provenance.append(["field", "value"])
    provenance.append(["software", f"SimNIBS {result['software']['version']}"])
    provenance.append(["schema", result["schema"]])
    readiness = wb.create_sheet("Figure readiness")
    readiness_columns = [
        "figure_id", "section", "scientific_question", "source", "machine_data",
        "unit", "coordinate_system", "status", "blocking_reason", "replacement_condition",
    ]
    readiness.append(readiness_columns)
    for item in figure_readiness:
        readiness.append([item.get(key) for key in readiness_columns])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="176B50")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(48, max(12, max(len(str(c.value or "")) for c in column) + 2))
            for cell in column: cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(TABLE / "inverse_ti_results.xlsx")


def main() -> None:
    result = json.loads((RAW / "inverse_results.json").read_text(encoding="utf-8"))
    validation = json.loads((OUT / "independent_fem" / "independent_validation.json").read_text(encoding="utf-8"))
    electrode_mesh_qc = json.loads((OUT / "independent_fem" / "electrode_mesh_qc.json").read_text(encoding="utf-8"))
    spatial_extrema = json.loads((FIG / "spatial_extrema.json").read_text(encoding="utf-8"))
    cluster_analysis = json.loads((TABLE / "four_electrode_offtarget_suprathreshold_clusters.json").read_text(encoding="utf-8"))
    cluster_sensitivity = json.loads((TABLE / "four_electrode_offtarget_suprathreshold_clusters_four_electrode_half_max_sensitivity.json").read_text(encoding="utf-8"))
    with np.load(OUT / "independent_fem" / "independent_timax_recompute.npz") as recompute:
        timax = recompute["TImax_V_per_m"]
        volumes = recompute["element_volumes_mm3"]
        target_mask = recompute["target_mask"]
        off_target_mask = recompute["off_target_mask"]
        gray_mask = recompute["whole_gray_matter_mask"]
        recompute_half_max_threshold = 0.5 * float(timax[gray_mask].max())

        def threshold_metrics(mask: np.ndarray) -> dict:
            selected = mask & (timax >= recompute_half_max_threshold)
            suprathreshold_volume = float(volumes[selected].sum())
            region_volume = float(volumes[mask].sum())
            return {
                "suprathreshold_tetrahedra": int(selected.sum()),
                "suprathreshold_volume_mm3": suprathreshold_volume,
                "threshold_coverage_pct": 100.0 * suprathreshold_volume / region_volume,
            }

        recompute_threshold_sensitivity = {
            "definition": "0.5 * four-active-electrode recomputation whole-gray-matter single-tetrahedron maximum",
            "timing": "post_hoc_sensitivity",
            "threshold_v_per_m": recompute_half_max_threshold,
            "target": threshold_metrics(target_mask),
            "off_target": threshold_metrics(off_target_mask),
        }
        sensitivity_total = (
            recompute_threshold_sensitivity["target"]["suprathreshold_volume_mm3"]
            + recompute_threshold_sensitivity["off_target"]["suprathreshold_volume_mm3"]
        )
        recompute_threshold_sensitivity["target_share_of_all_suprathreshold_gray_pct"] = (
            100.0 * recompute_threshold_sensitivity["target"]["suprathreshold_volume_mm3"] / sensitivity_total
        )
    validation["recompute_half_max_threshold_sensitivity"] = recompute_threshold_sensitivity
    for directory in (FIG, TABLE, METHOD, REPRO / "scripts", RAW / "anatomy"): directory.mkdir(parents=True, exist_ok=True)
    shutil.copy2(M2M / "T1.nii.gz", RAW / "anatomy" / "ernie_T1_conform.nii.gz")
    shutil.copy2(M2M / "final_tissues.nii.gz", RAW / "anatomy" / "ernie_final_tissues.nii.gz")
    shutil.copy2(M2M / "final_tissues_LUT.txt", RAW / "anatomy" / "ernie_final_tissues_LUT.txt")
    figure_python = SIMNIBS_PYTHON if SIMNIBS_PYTHON.is_file() else Path(sys.executable)
    subprocess.run(
        [str(figure_python), str(ROOT / "scripts" / "build_simnibs_inverse_publication_context_figures.py")],
        check=True,
        timeout=180,
    )
    reproduction_scripts = [
        "run_simnibs_ti_customer_demo.py",
        "run_simnibs_inverse_ti_discrete_ernie.py",
        "validate_simnibs_inverse_ti_winner.py",
        "plot_simnibs_inverse_ti_field.py",
        "build_simnibs_inverse_ti_delivery.py",
        "refresh_simnibs_inverse_manifest.py",
        "acceptance_simnibs_inverse_report_visual.mjs",
        "acceptance_simnibs_dual_delivery.py",
        "analyze_simnibs_inverse_offtarget_clusters.py",
        "export_simnibs_inverse_offtarget_mask_nifti.py",
        "export_simnibs_inverse_four_electrode_timax_nifti.py",
        "export_simnibs_inverse_final_timax_nifti.py",
        "audit_simnibs_inverse_electrode_mesh.py",
        "build_simnibs_inverse_publication_context_figures.py",
        "simnibs_inverse_report_content.py",
    ]
    for script_name in reproduction_scripts:
        shutil.copy2(ROOT / "scripts" / script_name, REPRO / "scripts" / script_name)
    (REPRO / "environment.json").write_text(json.dumps({
        "solver": {"name": "SimNIBS", "version": "4.6.0", "evidence": "raw/fem/simnibs_simulation_20260729-024122.log"},
        "platform": "Windows",
        "required_python_packages": ["simnibs==4.6.0", "numpy", "h5py", "nibabel", "openpyxl", "Pillow", "matplotlib"],
        "uncaptured_versions": ["Python", "NumPy", "h5py", "nibabel", "openpyxl", "Pillow", "matplotlib"],
        "note": "除 SimNIBS 外，其余运行时版本未在原始 FEM 运行时冻结；这属于逐位环境复现的限制，不得据此声称跨环境逐位一致。"
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    (REPRO / "README.md").write_text("""# 复算入口

1. 在项目根目录准备 `data/simnibs_examples_v4_1/extracted/m2m_ernie`，使用 SimNIBS 4.6.0 环境。
2. 运行 `scripts/run_simnibs_inverse_ti_discrete_ernie.py` 重建共享候选网格、求解三个基场并穷举候选；已有 FEM 时可加 `--reuse-fem`。
3. 运行 `scripts/validate_simnibs_inverse_ti_winner.py`，沿用同一 ernie 头组织网格，仅重建 F5、P5、FC3、CP3 四个活动电极及其邻近区域并重新求解排名第一方案。该步骤是电极构型场解复算，不是网格独立性验证。
4. 运行 `scripts/analyze_simnibs_inverse_offtarget_clusters.py` 以共享面邻接计算主阈值下的靶外超阈连通域；再运行 `scripts/analyze_simnibs_inverse_offtarget_clusters.py --threshold 0.3458170967431015 --suffix _four_electrode_half_max_sensitivity` 计算四活动电极场自身 50% 极值阈值下的敏感性聚类。运行 `scripts/export_simnibs_inverse_offtarget_mask_nifti.py` 生成主阈值的质心采样 NIfTI 定位掩膜。
5. 运行 `scripts/export_simnibs_inverse_four_electrode_timax_nifti.py`，把四活动电极最终 TImax 灰质四面体质心落到 ernie conform 体素；该 NIfTI 仅用于叠加定位，ROI 定量仍使用 Msh/NPZ。
6. 运行 `scripts/export_simnibs_inverse_final_timax_nifti.py`，用 SimNIBS 单元赋值把四活动电极最终灰质 TImax 导出到 ernie T1 1 mm 网格；该文件用于体空间叠加和论文再制图，四面体级 ROI 定量仍以 Msh/NPZ/JSON 为准。
7. 运行 `scripts/plot_simnibs_inverse_ti_field.py` 和 `scripts/build_simnibs_inverse_ti_delivery.py` 生成图表、工作簿与 HTML。
8. 运行视觉验收、刷新 manifest，再运行双模式验收。所有脚本按本项目目录结构解析输入；迁移到其他目录时应显式修改路径配置。

原始 FEM 日志是 SimNIBS 4.6.0 的版本证据。Python 及其余依赖的精确版本未在原始运行时冻结，因此本包支持算法和数据复算，但不承诺跨环境逐位一致。
""", encoding="utf-8")
    shutil.copy2(RAW / "tables" / "candidate_ranking.csv", TABLE / "candidate_ranking.csv")
    with (TABLE / "electrode_geometry.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["electrode_a", "electrode_b", "center_distance_mm", "nominal_edge_gap_mm", "definition"])
        for electrode_a, electrode_b, center_distance, edge_gap in ELECTRODE_GEOMETRY:
            writer.writerow([electrode_a, electrode_b, center_distance, edge_gap, "Euclidean center distance minus 40 mm nominal diameter"])
    (RAW / "fields" / "ernie_inverse_ti_TImax_discrete_optimum.json").write_text(
        json.dumps(
            {
                "file": "ernie_inverse_ti_TImax_discrete_optimum.nii.gz",
                "role": "search_field_candidate_ranking_only",
                "prohibited_use": "Do not use as the selected montage value-of-record, final-result statistics, or publication result figure.",
                "space": "ernie individual conform space",
                "units": "V/m",
                "generation": "SimNIBS transformations.interpolate_to_volume(method='linear', continuous=False, keep_tissues=[WM, GM])",
                "shape": [256, 256, 208],
                "voxel_size_mm": [1.0, 1.0, 1.0],
                "covered_wm_gm_voxels": 1322268,
                "reference_wm_gm_voxels": 1322268,
                "covered_wm_gm_voxel_pct": 100.0,
                "maximum_v_per_m": 0.7436366677284241,
                "maximum_scope_note": "Volume includes WM and GM; it is not the gray-matter-only single-tetrahedron maximum reported in the search table.",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    winner = result["winner"]
    candidates = result["all_candidates"]
    save_chart("figure_1_candidate_objective", "搜索阶段：候选方案的目标区剂量", "三个可行方案均由真实 SimNIBS FEM 基场组合计算", [c["candidate_id"].replace("__", " + ") for c in candidates], [c["target_mean_v_per_m"] for c in candidates], "V/m", "#245F8F")
    save_chart("figure_2_target_offtarget", "搜索阶段：最优方案的目标区与靶外灰质比较", f"有限集合最优：{winner['carrier_1']} + {winner['carrier_2']}；未由独立网格复核", ["目标 ROI 均值", "靶外灰质均值", "目标 ROI P95", "靶外灰质 P95"], [result["metrics"]["target"]["mean"], result["metrics"]["off_target"]["mean"], result["metrics"]["target"]["p95"], result["metrics"]["off_target"]["p95"]], "V/m", "#176B50")
    save_chart("figure_3_finite_search", "有限搜索空间的目标函数排序", "第 2、3 名仅差约 0.5%，当前证据下不可区分；本图不是收敛曲线", ["Rank 1", "Rank 2/3", "Rank 2/3"] , [c["objective_value_v_per_m"] for c in candidates], "V/m", "#A36B18")
    save_chart("figure_4_constraints", "固定电流输入核验", "每个载波回路固定为 +1/-1 mA 峰值，净电流为 0；未进行幅度优化", ["回路 1 正电流", "回路 1 负电流绝对值", "回路 2 正电流", "回路 2 负电流绝对值"], [1, 1, 1, 1], "mA peak", "#A1433F")
    off_target_delta_pct = 100.0 * (validation["independent_metrics"]["off_target"]["mean"] - result["metrics"]["off_target"]["mean"]) / result["metrics"]["off_target"]["mean"]
    mean_ratio_delta_pct = 100.0 * (validation["independent_target_to_off_target_mean_ratio"] - winner["target_to_off_target_mean_ratio"]) / winner["target_to_off_target_mean_ratio"]
    save_chart("figure_5_independent_fem", "搜索场与四活动电极复算场的描述性对照", f"ROI mean +{validation['relative_difference_pct']:.2f}%；靶外 mean {off_target_delta_pct:+.2f}%；均值比 {mean_ratio_delta_pct:+.2f}%；差异不可归因", ["搜索目标 mean", "搜索靶外 mean", "复算目标 mean", "复算靶外 mean"], [result["metrics"]["target"]["mean"], result["metrics"]["off_target"]["mean"], validation["independent_metrics"]["target"]["mean"], validation["independent_metrics"]["off_target"]["mean"]], "V/m", "#6B5B95")
    publication_figures = {
        "figure_0_optimized_timax_slices": "fig01_optimized_timax_slices",
        "figure_1_candidate_objective": "fig02_candidate_objective",
        "figure_2_target_offtarget": "fig03_target_offtarget",
        "figure_3_finite_search": "fig04_candidate_ranking",
        "figure_4_constraints": "fig05_current_constraints",
        "figure_5_independent_fem": "fig06_independent_fem",
        "figure_6_four_electrode_timax_slices": "fig07_four_electrode_timax_slices",
        "figure_7_four_electrode_quantitative": "fig08_four_electrode_quantitative",
    }
    for source, destination in publication_figures.items():
        for extension in ("png", "svg"):
            shutil.copy2(FIG / f"{source}.{extension}", FIG / f"{destination}.{extension}")
    figure_readiness = [
        {"figure_id": "model_target_registration", "section": "primary_result", "scientific_question": "头模型、组织边界和目标 ROI 在个体空间中的位置是什么？", "source": "figures/fig00_model_target_registration.png", "machine_data": "raw/anatomy/ernie_T1_conform.nii.gz; raw/anatomy/ernie_final_tissues.nii.gz", "unit": "mm", "coordinate_system": "ernie individual conform physical space", "status": "demo_only", "blocking_reason": "示例受试者，未使用客户 MRI；ROI 未完成 atlas/MNI 验证", "replacement_condition": "客户 MRI 分割、配准与靶点审核完成后重绘"},
        {"figure_id": "stimulation_montage", "section": "primary_result", "scientific_question": "最终四电极方案的空间位置、极性和载波回路如何？", "source": "figures/fig00_inverse_electrode_montage.png", "machine_data": "figures/figure_data/fig00_inverse_electrode_montage.csv", "unit": "mA peak; mm", "coordinate_system": "ernie individual conform physical space", "status": "demo_only", "blocking_reason": "F5-FC3 存在两节点点接触；未映射客户 64 通道帽", "replacement_condition": "按真实触点几何验证正间隙并重新求解"},
        {"figure_id": "carrier_and_timax_field", "section": "primary_result", "scientific_question": "最终 E1、E2 和 TImax 在目标高度附近如何分布？", "source": "figures/fig00_four_electrode_e1_e2_timax.png", "machine_data": "independent_fem/independent_timax_recompute.npz", "unit": "V/m", "coordinate_system": "ernie individual conform physical space", "status": "demo_only", "blocking_reason": "灰质轴向 slab；点接触网格尚未替换", "replacement_condition": "正式电极网格重算并按预定色标重绘"},
        {"figure_id": "field_surface", "section": "primary_result", "scientific_question": "最终 TImax 在灰质表面如何分布？", "source": "figures/fig00_four_electrode_timax_surface.png", "machine_data": "figures/figure_data/fig00_final_surface_context.npz", "unit": "V/m", "coordinate_system": "ernie individual conform physical space", "status": "demo_only", "blocking_reason": "点接触网格；表面值为灰质表面插值", "replacement_condition": "正式网格重算并复核表面插值"},
        {"figure_id": "four_electrode_field", "section": "primary_result", "scientific_question": "最终 TImax 在 ROI 切面及靶外极值切面如何分布？", "source": "figures/fig07_four_electrode_timax_slices.png", "machine_data": "independent_fem/independent_timax_recompute.npz", "unit": "V/m", "coordinate_system": "ernie individual conform physical space", "status": "demo_only", "blocking_reason": "点接触网格；极值位置未做稳健性分析", "replacement_condition": "正式网格和稳健性分析后重绘"},
        {"figure_id": "four_electrode_quantitative", "section": "primary_result", "scientific_question": "目标区与靶外灰质的场强、覆盖率和超阈体积有何差异？", "source": "figures/fig08_four_electrode_quantitative.png", "machine_data": "report_data.json / results.roi_metrics; tables/inverse_ti_results.xlsx / Final four-electrode metrics", "unit": "V/m; %; mm3", "coordinate_system": "regional statistics", "status": "demo_only", "blocking_reason": "阈值为事后探索性；点接触网格", "replacement_condition": "正式研究预注册阈值或多阈值方案并重新求解"},
        {"figure_id": "candidate_ranking", "section": "optimization_evidence", "scientific_question": "有限候选库内哪个组合最大化目标函数？", "source": "figures/fig02_candidate_objective.png", "machine_data": "tables/candidate_ranking.csv", "unit": "V/m", "coordinate_system": "ROI statistic", "status": "exploratory_search_only", "blocking_reason": "共享搜索网格含未激活电极材料；只有入选方案完成四电极复算", "replacement_condition": "正式搜索网格重建并复算所有待比较方案"},
        {"figure_id": "finite_search_order", "section": "optimization_evidence", "scientific_question": "有限候选库的完整名义排序是什么？", "source": "figures/fig04_candidate_ranking.png", "machine_data": "tables/candidate_ranking.csv", "unit": "V/m", "coordinate_system": "ROI statistic", "status": "exploratory_search_only", "blocking_reason": "第 2、3 名未完成四活动电极构型复算和排序不确定度分析", "replacement_condition": "正式搜索网格重建并复算全部待比较方案"},
        {"figure_id": "current_constraints", "section": "optimization_evidence", "scientific_question": "两个载波回路采用了什么固定电流约束？", "source": "figures/fig05_current_constraints.png", "machine_data": "report_data.json / protocol.electrodes", "unit": "mA peak", "coordinate_system": "circuit current", "status": "method_evidence", "blocking_reason": "未优化载波幅度分配，未指定频率和相位", "replacement_condition": "正式设备协议确定后更新电流、频率和相位"},
        {"figure_id": "four_electrode_recompute", "section": "optimization_evidence", "scientific_question": "搜索场与四活动电极复算场的区域均值有何差异？", "source": "figures/fig06_independent_fem.png", "machine_data": "independent_fem/independent_validation.json", "unit": "V/m", "coordinate_system": "regional statistics", "status": "demo_only", "blocking_reason": "搜索阶段缺少同口径校准误差汇总，跨阶段差异不能归因于电极构型，也不是网格收敛或稳健性证据", "replacement_condition": "以同口径校准 QC 完成正式几何重算，并执行网格收敛分析"},
    ]
    electrical_validity_block = (
        "F5-FC3 两节点点接触使两个载波回路的电气独立性未建立；局部分流可能影响 "
        "E1、E2、TImax 及派生统计，当前数值和图件不得作为正式研究结果引用"
    )
    electrically_affected_figures = {
        "stimulation_montage",
        "carrier_and_timax_field",
        "field_surface",
        "four_electrode_field",
        "four_electrode_quantitative",
        "candidate_ranking",
        "finite_search_order",
        "four_electrode_recompute",
    }
    for figure in figure_readiness:
        if figure["figure_id"] in electrically_affected_figures:
            figure["status"] = "demo_only"
            figure["blocking_reason"] = f"{electrical_validity_block}；{figure['blocking_reason']}"
            figure["replacement_condition"] = (
                "验证正的最小电极间隙，重建电极网格，重新求解两个载波场并替换全部受影响图件和统计"
            )
    workbook(result, validation, figure_readiness)

    conductivities_text = "；".join(f"{name} {value:.3f}" for name, value in CONDUCTIVITIES_S_PER_M.items())
    method_text = f"""# 方法与解释边界

本演示使用 SimNIBS {result['software']['version']} 官方 ernie 示例头模型。搜索空间由三个双极电极对组成，程序在同一电极化有限元网格上分别求解 1 mA 峰值基场，再使用 `simnibs.utils.TI_utils.get_maxTI` 计算所有互不重叠的双载波组合。目标函数为 20 mm 球形灰质 ROI 内 TImax 的四面体体积加权均值。选中方案随后仅使用四个活动电极重建模型并复算，ROI 均值相对差异为 {validation['relative_difference_pct']:.2f}%。本演示没有可核验的事前接受阈值，因此该差异只作描述。四活动电极求解日志的估计电流校准误差最高为 6.9%，与 6.25% 的场解差异同量级；搜索场又缺少同口径的校准误差汇总。两套结果使用相同的 ROI 灰质单元和体积权重，这只说明 ROI 定义、插值位置和统计权重一致。由于未完成网格收敛测试，该差异无法可靠分解为电极几何、场解变化和数值离散误差，也不能作为网格误差上限或 ROI 数值稳健性证据。

搜索阶段电极在 SimNIBS 中定义为 `shape="ellipse"`、`dimensions=[40, 40]`，即长短轴均为 40 mm 的圆形电极，直径 40 mm、厚 2 mm，几何面积约 1256.6 mm²（12.566 cm²）。每个活动回路输入 1 mA peak，据几何面积折算的名义平均输入密度为 0.0796 mA/cm²；它不是有限元接触面的局部电流密度，本次未统计局部电流密度，不能用于安全限值判断。电极化网格含 {SEARCH_MESH['nodes']:,} 个节点、{SEARCH_MESH['tetrahedra']:,} 个四面体和 {SEARCH_MESH['triangles']:,} 个三角形。ROI 归属按四面体重心判定：组织标签为灰质，且重心到个体空间坐标 (-41, -13, 66) mm 的欧氏距离不超过 20 mm。各向同性组织电导率（S/m）为：{conductivities_text}。

四活动电极场解复算沿用同一 ernie 头组织网格，仅保留 F5、P5、FC3、CP3 四个活动电极，并重新构建电极及其邻近区域后重新求解；活动电极尺寸与组织电导率不变。四电极复算网格含 {INDEPENDENT_MESH['nodes']:,} 个节点、{INDEPENDENT_MESH['tetrahedra']:,} 个四面体和 {INDEPENDENT_MESH['triangles']:,} 个三角形。参与灰质统计的单元、中心和体积权重与搜索场逐元素一致，因此这不是网格独立性或灰质离散误差验证。搜索网格还包含候选库中的非活动电极，四活动电极复算同时重建了电极及邻近区域；复算日志记录 0.3%–6.9% 的估计电流校准误差，而搜索阶段缺少同口径汇总。因此，观测到的 6.25% 跨阶段差异不能与运行间校准偏差分离，也不能归因于电极构型。灰质网格误差仍需单独的网格收敛测试。

TImax 使用 SimNIBS 4.6.0 的 `simnibs.utils.TI_utils.get_maxTI(E1, E2)`。算法逐单元交换矢量使 ||E1||≥||E2||，必要时翻转 E2 使夹角 α≤90°；若 ||E2||≤||E1||cosα，则 TImax=2||E2||，否则 TImax=2||E2×(E1−E2)||/||E1−E2||。因此本报告数值包含公式中的因子 2，表示沿所有方向取得的最大调制幅值，不是 ||E2|| 半幅、RMS 或载波峰峰值。E1、E2 均按每回路 1 mA peak 标定；在线性准静态假设下，两个回路电流同时按相同比例缩放时，TImax 按同一比例缩放。若设备或文献使用 RMS、峰峰值或其他电流口径，必须先换算到本报告的峰值电流与最大调制幅值口径。有限元求解采用准静态、线性场叠加假设；本演示未指定载波频率 f1、f2 或差频 Δf，频率没有参与电场或 TImax 数值计算。因此这些结果不能用于比较频率效应。正式刺激协议必须另行声明 f1、f2、Δf、相位和设备波形约定。

超阈体积采用固定绝对阈值 {result['metrics']['target']['threshold_v_per_m']:.6f} V/m。该数值由搜索场全灰质单四面体最大值的 50% 一次性导出，并原样用于搜索网格和独立网格；它不是各网格分别计算的 50% 相对阈值。该阈值为事后探索性定义，不是预注册阈值，也受搜索网格单元极值影响；超阈覆盖仅用于本演示内部描述。正式研究应在分析前声明阈值，或报告多阈值敏感性分析。

区域 mean、median、P05 和 P95 均按四面体体积加权；加权分位数将单元按 TImax 升序排列，在累计体积首次达到区域总体积的 5%、50% 或 95% 时取对应单元值，不做插值。max 是区域内未加权的单四面体极值。超阈体积为满足 TImax ≥ 阈值的四面体体积之和，覆盖率分母为同一区域的总体积。

候选集合有限且已穷举，因此“集合内最优”可以核验；迭代收敛曲线不适用。结果不能解释为连续头皮位置的全局最优，也没有完成客户 64 通道帽的离散映射。本场解使用直径 40 mm、厚 2 mm 的圆形电极；帽式小触点即使使用同名 10-10 位置，也必须按真实触点尺寸、接触层和位置重新求解，不能直接迁移场强、局部电流密度或暴露结论。原生 `TesFlexOptimization` 尝试在头模型准备阶段发生 Windows 原生异常，未产生优化结果，本交付包没有引用该失败任务的数值。

共享搜索网格包含候选库全部电极几何；每次基场求解时，未激活电极使用零电流通道，但仍作为电极材料参与有限元模型，可能改变分流和局部电场。共享几何用于保证三个基场网格和单元顺序完全一致，不代表真实四活动电极构型。当前仅对排名第一的组合执行了四活动电极场解复算，因此搜索阶段绝对场强以及未复算候选的名义排名都不能解释为设备构型下已验证的结果。该复算不是网格独立性验证。

本次定量分析域仅包含灰质；CSF、颅骨、头皮等非灰质组织暴露评估未执行，属于正式研究交付前的硬门禁。电场结果只描述已分析组织中的模型场分布，不等同于神经激活、全组织暴露评估、临床安全性、治疗效果或个体处方。
"""
    (METHOD / "methods_and_boundaries.md").write_text(method_text, encoding="utf-8")
    bib_entries = []
    for reference in REFERENCES:
        bib_entries.append(
            "@misc{{{id},\n  author = {{{authors}}},\n  title = {{{title}}},\n  year = {{{year}}},\n"
            "  doi = {{{doi}}},\n  note = {{PMID: {pmid}}}\n}}".format(**reference)
        )
    (METHOD / "references.bib").write_text("\n\n".join(bib_entries) + "\n", encoding="utf-8")
    with (METHOD / "software_versions.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["software", "version", "purpose", "evidence"])
        writer.writerow(["SimNIBS", result["software"]["version"], "FEM solve and TImax calculation", "raw/fem/simnibs_simulation_20260729-024122.log"])
        writer.writerow(["Microsoft Edge", "runtime executable; exact version not frozen", "HTML to PDF rendering", "report.pdf build step"])
    run_id = "ernie-inverse-ti-discrete-three-candidate-20260729"
    (METHOD / "data_dictionary.md").write_text("""# 数据字典

- `report_data.json / model.electrode_material_model`: 实际 FEM 使用的单层 2 mm `Electrode_rubber` 电极及 29.4 S/m 电导率；未建模凝胶、海绵或接触阻抗。
- `report_data.json / results.roi_metrics`: 四活动电极复算场的最终引用口径。
- `report_data.json / results.search_stage_roi_metrics`: 共享候选网格的搜索阶段指标，仅用于候选排序。
- `report_data.json / results.cross_stage_comparison`: 入选方案从搜索场到四活动电极复算场的 ROI mean、靶外 mean 与目标/靶外均值比变化；仅作跨阶段描述性比较。搜索阶段缺少同口径校准误差汇总，不能分离运行间校准偏差，也不能归因于电极构型。
- `report_data.json / results.recompute_half_max_threshold_sensitivity`: 四活动电极复算场自身 50% 全灰质单四面体极值阈值下的敏感性结果。
- Excel `Final four-electrode metrics`: 论文引用与二次统计的最终 ROI 指标。
- Excel `Search-stage metrics`: 搜索阶段指标，仅用于候选排序。
- Excel `Search-final comparison`: 三项搜索场到四活动电极复算场的对照值、相对变化和解释边界。
- Excel `Figure readiness`: 每张图的科学问题、来源、机器数据和正式研究阻断状态。
- `candidate_ranking.csv`: 每行一个可行双载波候选。
- `candidate_timax_fields.npz`: 各候选逐单元 TImax，单位 V/m。
- `inverse_ti_results.h5`: 单元中心、四面体体积、ROI 掩膜与全部候选搜索场。
- `ernie_inverse_ti_optimum.msh`: 共享候选库网格上的最优候选 TImax 搜索场。
- `raw/anatomy/ernie_T1_conform.nii.gz`: T1grid 的解剖参考影像；与主要体空间文件具有完全相同的 shape、1 mm 体素和 affine，可直接叠加，无需额外配准变换。
- `raw/anatomy/ernie_final_tissues.nii.gz` 与 LUT: ernie 个体空间组织标签及标签字典，用于复核头模型和 ROI 的解剖语境。
- `independent_fem/independent_timax_recompute.npz`: 最终结果逐单元数组，键为 `E1_V_per_m`、`E2_V_per_m`、`TImax_V_per_m`、`element_centers_mm`、`element_volumes_mm3`、`target_mask`、`off_target_mask`、`whole_gray_matter_mask`。
- `figures/figure_data/fig00_final_surface_context.npz`: 灰质/头皮表面节点、三角形与最终 TImax 表面插值值。
- 带 `_four_electrode_half_max_sensitivity` 后缀的 CSV/JSON: 阈值敏感性下的靶外连通域结果。
- `four_electrode_timax_T1grid.nii.gz`: 四活动电极最终 TImax 灰质场的主要体空间文件；SimNIBS 四面体单元赋值到 ernie T1 1 mm conform 网格，不平滑，非灰质置零。用于 FSL/SPM/nilearn 叠加和再制图，不替代 Msh/NPZ 四面体级定量，也不表示非灰质暴露。具体 affine 与参考影像哈希见同名 JSON。
- `four_electrode_timax_conform_nearest_centroid.nii.gz`: 质心落栅定位辅助文件；同体素取最大值，未采样体素为 0，仅覆盖 70.37% 的参考灰质体素，不可作为主要体场或定量输入。
- `TImax`: SimNIBS 4.6.0 `get_maxTI` 的最大调制幅值；公式包含因子 2，不是半幅、RMS 或载波峰峰值。E1/E2 按每回路 1 mA peak 标定。
- `mean`, `median`, `p05`, `p95`: 均按四面体体积加权；分位数按 TImax 升序累计体积并取首次达到目标概率的单元值，不插值。
- `max`: 区域内未加权的单四面体极值。
- `suprathreshold_volume_mm3`: TImax 大于等于阈值的四面体体积之和。
- `threshold_coverage_pct`: 100 × 超阈体积 / 同一区域总体积。
""", encoding="utf-8")
    roi_metrics = []
    search_stage_roi_metrics = []
    for key, role, label in (("target", "目标区", "20 mm 球形灰质 ROI"), ("off_target", "非目标区", "目标 ROI 外全部灰质"), ("whole_gray_matter", "分析域参考", "全部灰质")):
        roi_metrics.append({"roi": label, "role": role, "definition": label, **validation["independent_metrics"][key], "run_id": validation["run_id"]})
        search_stage_roi_metrics.append({"roi": label, "role": role, "definition": label, **result["metrics"][key], "run_id": run_id})
    artifact_specs = [
        ("raw/fields/ernie_inverse_ti_optimum.msh", "MSH", "共享候选库电极化网格上的最优 TImax 搜索场"),
        ("raw/fields/ernie_inverse_ti_TImax_discrete_optimum.nii.gz", "NIFTI", "仅用于候选排序场诊断；禁止作为最终方案统计或论文结果图"),
        ("raw/fields/ernie_inverse_ti_TImax_discrete_optimum.json", "JSON", "搜索场 NIfTI 的生成方式、分辨率、覆盖率和禁止用途"),
        ("raw/anatomy/ernie_T1_conform.nii.gz", "NIFTI", "四活动电极 T1grid 的 ernie 个体 conform 解剖参考影像；同 shape、体素和 affine"),
        ("raw/anatomy/ernie_final_tissues.nii.gz", "NIFTI", "ernie 个体空间组织标签，用于头模型和 ROI 解剖语境复核"),
        ("raw/anatomy/ernie_final_tissues_LUT.txt", "TEXT", "final_tissues 组织标签字典"),
        ("raw/fields/inverse_ti_results.h5", "HDF5", "共享候选库网格的全部候选场、ROI 掩膜和网格量"),
        ("tables/candidate_ranking.csv", "CSV", "候选方案完整排名"),
        ("tables/electrode_geometry.csv", "CSV", "ernie 10-10 电极中心距与 40 mm 名义直径边缘间隙"),
        ("tables/inverse_ti_results.xlsx", "EXCEL", "二次统计工作簿"),
        ("raw/fields/candidate_timax_fields.npz", "NPZ", "全部候选逐单元 TImax 数组"),
        ("independent_fem/independent_validation.json", "JSON", "入选方案四活动电极构型复算结果"),
        ("independent_fem/independent_timax_recompute.npz", "NPZ", "四活动电极构型复算的 E1、E2、TImax、网格量与目标/靶外掩膜"),
        ("independent_fem/four_electrode_timax_T1grid.nii.gz", "NIFTI", "四活动电极最终 TImax 灰质场的主要体空间文件；SimNIBS 单元赋值到 T1 1 mm 网格，非灰质置零"),
        ("independent_fem/four_electrode_timax_T1grid.json", "JSON", "主要体空间 TImax 的来源、空间、单元赋值方法和定量限制"),
        ("independent_fem/four_electrode_timax_conform_nearest_centroid.nii.gz", "NIFTI", "四活动电极最终 TImax 灰质场的 conform 质心落栅体数据，用于叠加定位"),
        ("independent_fem/four_electrode_timax_conform_nearest_centroid.json", "JSON", "质心落栅 NIfTI 的来源、空间、70.37% 灰质体素覆盖和定量限制"),
        ("independent_fem/electrode_mesh_qc.json", "JSON", "四活动电极体的共享节点/面/四面体、最小节点距和退化单元核验"),
        ("raw/fem/simnibs_simulation_20260729-024122.mat", "MAT", "共享候选网格基场求解的 SimNIBS 会话参数"),
        ("independent_fem/simnibs_simulation_20260729-030053.mat", "MAT", "四活动电极构型复算的 SimNIBS 会话参数"),
        ("independent_fem/simnibs_simulation_20260729-030053.log", "LOG", "四活动电极构型复算日志及估计电流校准误差证据"),
        ("independent_fem/ernie_TDCS_1_scalar.msh", "MSH", "四活动电极独立复算的载波 1 矢量电场 E；scalar 是 SimNIBS 输出文件名，不表示仅含幅值"),
        ("independent_fem/ernie_TDCS_2_scalar.msh", "MSH", "四活动电极独立复算的载波 2 矢量电场 E；scalar 是 SimNIBS 输出文件名，不表示仅含幅值"),
        ("methods/methods_and_boundaries.md", "MARKDOWN", "计算方法与解释边界"),
        ("methods/data_dictionary.md", "MARKDOWN", "二次分析数据字典"),
        ("methods/references.bib", "BIBTEX", "经 PubMed 核验的软件与 TI 方法引用"),
        ("methods/software_versions.csv", "CSV", "软件版本、用途和证据来源"),
        ("reproduction/README.md", "MARKDOWN", "从真实 FEM 到报告的复算顺序、输入约定与环境限制"),
        ("reproduction/environment.json", "JSON", "已捕获的软件版本、依赖要求及未冻结版本"),
        ("reproduction/scripts/run_simnibs_inverse_ti_discrete_ernie.py", "PYTHON", "共享候选网格真实 FEM 与有限候选穷举"),
        ("reproduction/scripts/validate_simnibs_inverse_ti_winner.py", "PYTHON", "排名第一方案的四活动电极场解复算"),
        ("reproduction/scripts/plot_simnibs_inverse_ti_field.py", "PYTHON", "TImax 空间场制图"),
        ("reproduction/scripts/build_simnibs_inverse_publication_context_figures.py", "PYTHON", "解剖、蒙太奇、载波场和灰质表面出版上下文图"),
        ("reproduction/scripts/simnibs_inverse_report_content.py", "PYTHON", "客户报告信息架构和 HTML 内容渲染"),
        ("reproduction/scripts/analyze_simnibs_inverse_offtarget_clusters.py", "PYTHON", "四活动电极复算场靶外超阈共享面连通域分析"),
        ("reproduction/scripts/export_simnibs_inverse_offtarget_mask_nifti.py", "PYTHON", "靶外超阈四面体质心采样 NIfTI 掩膜导出"),
        ("reproduction/scripts/export_simnibs_inverse_four_electrode_timax_nifti.py", "PYTHON", "四活动电极最终 TImax 灰质场的 conform 质心落栅 NIfTI 导出"),
        ("reproduction/scripts/export_simnibs_inverse_final_timax_nifti.py", "PYTHON", "四活动电极最终 TImax 灰质场的 SimNIBS 单元赋值 T1 网格 NIfTI 导出"),
        ("reproduction/scripts/build_simnibs_inverse_ti_delivery.py", "PYTHON", "客户 HTML、图表、Excel 与数据索引构建"),
        ("figures/fig01_optimized_timax_slices.png", "PNG", "共享候选网格搜索场诊断图；不作为最终方案结果图"),
        ("figures/fig01_optimized_timax_slices.svg", "SVG", "共享候选网格搜索场诊断矢量图；不作为最终方案结果图"),
        ("figures/fig02_candidate_objective.png", "PNG", "有限候选库目标函数比较"),
        ("figures/fig02_candidate_objective.svg", "SVG", "有限候选库目标函数比较矢量图"),
        ("figures/fig03_target_offtarget.png", "PNG", "搜索阶段目标区与靶外灰质诊断比较；不作为最终方案结果图"),
        ("figures/fig03_target_offtarget.svg", "SVG", "搜索阶段目标区与靶外灰质诊断比较矢量图"),
        ("figures/fig04_candidate_ranking.png", "PNG", "有限候选库完整排序"),
        ("figures/fig04_candidate_ranking.svg", "SVG", "有限候选库完整排序矢量图"),
        ("figures/fig05_current_constraints.png", "PNG", "两个载波回路的固定峰值电流约束"),
        ("figures/fig05_current_constraints.svg", "SVG", "两个载波回路的固定峰值电流约束矢量图"),
        ("figures/fig06_independent_fem.png", "PNG", "搜索阶段与四活动电极复算阶段的区域均值描述性对照；差异不可归因"),
        ("figures/fig06_independent_fem.svg", "SVG", "搜索阶段与四活动电极复算阶段的区域均值描述性对照矢量图；差异不可归因"),
        ("figures/fig07_four_electrode_timax_slices.png", "PNG", "四活动电极演示性复算空间场；受 F5-FC3 点接触限制"),
        ("figures/fig07_four_electrode_timax_slices.svg", "SVG", "四活动电极演示性复算空间场矢量图；正式制图须重新求解"),
        ("figures/fig08_four_electrode_quantitative.png", "PNG", "四活动电极演示性定量结果；受 F5-FC3 点接触限制"),
        ("figures/fig08_four_electrode_quantitative.svg", "SVG", "四活动电极演示性定量结果矢量图；正式制图须重新求解"),
        ("figures/fig00_model_target_registration.png", "PNG", "个体 T1、组织边界与目标 ROI 三正交定位"),
        ("figures/fig00_model_target_registration.svg", "SVG", "头模型与目标定位矢量图"),
        ("figures/fig00_inverse_electrode_montage.png", "PNG", "最终四电极位置、极性与载波回路"),
        ("figures/fig00_inverse_electrode_montage.svg", "SVG", "最终四电极蒙太奇矢量图"),
        ("figures/fig00_four_electrode_e1_e2_timax.png", "PNG", "最终 E1、E2 与 TImax 共色标灰质轴向场"),
        ("figures/fig00_four_electrode_e1_e2_timax.svg", "SVG", "最终载波场与 TImax 矢量图"),
        ("figures/fig00_four_electrode_timax_surface.png", "PNG", "最终 TImax 灰质表面插值分布"),
        ("figures/fig00_four_electrode_timax_surface.svg", "SVG", "最终 TImax 灰质表面矢量图"),
        ("figures/figure_data/fig00_final_surface_context.npz", "NPZ", "灰质/头皮表面节点、三角形与最终表面 TImax"),
        ("figures/figure_data/fig00_inverse_electrode_montage.csv", "CSV", "四电极坐标、电流极性与载波归属"),
        ("figures/figure_data/fig00_publication_context.json", "JSON", "出版上下文图的单位、色标和生成语义"),
        ("figures/spatial_extrema.json", "JSON", "搜索场与四活动电极复算场的靶外极值坐标及 ROI 中心距离"),
        ("tables/four_electrode_offtarget_suprathreshold_clusters.csv", "CSV", "四活动电极复算场靶外超阈连通域表"),
        ("tables/four_electrode_offtarget_suprathreshold_clusters.json", "JSON", "靶外超阈连通域定义、完整精度指标与质心坐标"),
        ("tables/four_electrode_offtarget_suprathreshold_clusters_four_electrode_half_max_sensitivity.csv", "CSV", "四活动电极场自身 50% 极值阈值下的靶外聚类敏感性表"),
        ("tables/four_electrode_offtarget_suprathreshold_clusters_four_electrode_half_max_sensitivity.json", "JSON", "四活动电极场自身 50% 极值阈值下的聚类数、主团体积和空间范围"),
        ("independent_fem/four_electrode_offtarget_suprathreshold_centers_mm_four_electrode_half_max_sensitivity.npy", "NPY", "四活动电极场自身 50% 极值阈值下的靶外超阈四面体中心坐标"),
        ("independent_fem/four_electrode_offtarget_suprathreshold_centroid_mask.nii.gz", "NIFTI", "靶外超阈四面体质心采样定位掩膜；不是体积栅格化"),
        ("independent_fem/four_electrode_offtarget_suprathreshold_centroid_mask.json", "JSON", "质心采样 NIfTI 的语义、参考空间和采样计数"),
        ("independent_fem/four_electrode_offtarget_suprathreshold_centers_mm.npy", "NPY", "靶外超阈四面体中心坐标"),
        ("independent_fem/four_electrode_offtarget_suprathreshold_mask.msh", "MSH", "靶外超阈四面体二值单元场"),
    ]
    fem_logs = sorted((RAW / "fem").glob("simnibs_simulation_*.log"))
    if fem_logs:
        artifact_specs.append((fem_logs[-1].relative_to(OUT).as_posix(), "LOG", "候选基场 FEM 求解日志"))
    recompute_figure_prefixes = (
        "figures/fig00_",
        "figures/fig06_",
        "figures/fig07_",
        "figures/fig08_",
        "figures/figure_5_independent_fem",
        "figures/figure_6_four_electrode_timax_slices",
        "figures/figure_7_four_electrode_quantitative",
        "figures/figure_data/fig00_",
    )

    def artifact_run_id(path: str) -> str:
        if (
            path.startswith("independent_fem/")
            or path.startswith("tables/four_electrode_")
            or path.startswith(recompute_figure_prefixes)
        ):
            return validation["run_id"]
        return run_id

    artifacts = [
        {
            "path": path,
            "type": kind,
            "purpose": purpose,
            "run_id": artifact_run_id(path),
            "sha256": sha256(OUT / path),
        }
        for path, kind, purpose in artifact_specs
    ]
    publication_context = json.loads((FIG / "figure_data" / "fig00_publication_context.json").read_text(encoding="utf-8"))
    carrier_display_vmax = publication_context["carrier_field_figure"]["shared_display_vmax_v_per_m"]
    surface_display_vmax = publication_context["surface_field_figure"]["display_vmax_v_per_m"]
    if carrier_display_vmax <= 0 or surface_display_vmax <= 0:
        raise ValueError("Publication-context figure color-scale maxima must be positive")
    whole_gray = validation["independent_metrics"]["whole_gray_matter"]
    threshold = validation["comparison_threshold_v_per_m"]
    search_roi_mean = validation["search_stage_roi_mean_v_per_m"]
    final_roi_mean = validation["independent_rerun_roi_mean_v_per_m"]
    search_off_target_mean = result["metrics"]["off_target"]["mean"]
    final_off_target_mean = validation["independent_metrics"]["off_target"]["mean"]
    search_mean_ratio = winner["target_to_off_target_mean_ratio"]
    final_mean_ratio = validation["independent_target_to_off_target_mean_ratio"]
    cross_stage_comparison = {
        "scope": "selected_montage_search_field_vs_four_active_electrode_recomputation",
        "comparison_type": "cross_stage_descriptive",
        "causal_attribution_status": "not_identifiable",
        "calibration_error_comparability": "not_comparable_search_stage_summary_missing",
        "configuration_effect_inference_allowed": False,
        "robustness_inference_allowed": False,
        "confounding_factors": [
            "electrodeized model changed between stages",
            "four-active-electrode recomputation log reports 0.3%-6.9% estimated current calibration error",
            "search-stage calibration-error summary is unavailable",
        ],
        "interpretation": "descriptive cross-stage difference only; run-to-run calibration deviation cannot be separated and the difference cannot be attributed to electrode configuration; not mesh convergence or robustness evidence",
        "metrics": [
            {"metric": "target_roi_mean_v_per_m", "search_stage": search_roi_mean, "four_electrode_recompute": final_roi_mean, "relative_change_pct": 100.0 * (final_roi_mean - search_roi_mean) / search_roi_mean},
            {"metric": "off_target_mean_v_per_m", "search_stage": search_off_target_mean, "four_electrode_recompute": final_off_target_mean, "relative_change_pct": 100.0 * (final_off_target_mean - search_off_target_mean) / search_off_target_mean},
            {"metric": "target_to_off_target_mean_ratio", "search_stage": search_mean_ratio, "four_electrode_recompute": final_mean_ratio, "relative_change_pct": 100.0 * (final_mean_ratio - search_mean_ratio) / search_mean_ratio},
        ],
    }
    report_data = {
        "schema_version": "simnibs.report.v1",
        "analysis_mode": "inverse_optimization",
        "stimulation_modality": "temporal_interference",
        "report": {"report_id": "QLA-SIMNIBS-INVERSE-TI-ERNIE-20260729-01", "title": "个体化逆向 TI 仿真报告", "status": "ready", "status_scope": "demo_delivery_package_generation", "publication_readiness": "formal_study_prerequisites_pending", "build_identity": run_id},
        "subject": {"id": "ernie（SimNIBS 官方示例受试者）"},
        "target": {"name": "左侧运动区演示 ROI", "definition": "ernie 个体空间 20 mm 球形灰质 ROI", "mni_coordinate_mm": None, "reference_example_mni_coordinate_mm": [-37.0, -21.0, 58.0], "reference_coordinate_status": "not_transformed_or_validated_in_this_delivery", "subject_coordinate_mm": result["target"]["center_mm"], "radius_mm": result["target"]["radius_mm"], "tissue": "gray matter", "coordinate_transform": "本次交付只使用 ernie 个体 conform 坐标；未执行或验证个体到 MNI 的空间变换。reference_example_mni_coordinate_mm 仅记录上游示例参考值，不是本报告验证后的 MNI 靶点。"},
        "model": {"head_model": "SimNIBS ernie", "solver": "SimNIBS 4.6.0 FEM", "solver_current_calibration_error_pct": {"minimum": 0.3, "maximum": 6.9, "source": "independent_fem/simnibs_simulation_20260729-030053.log", "posthoc_rescaling": False}, "conductivity_model": "isotropic", "conductivities_s_per_m": CONDUCTIVITIES_S_PER_M, "search_mesh": SEARCH_MESH, "independent_four_active_electrode_mesh": INDEPENDENT_MESH, "electrode_geometry": {"shape": "ellipse", "dimensions_mm": [40.0, 40.0], "thickness_mm": 2.0}, "electrode_material_model": ELECTRODE_MATERIAL_MODEL, "electrode_mesh_qc": electrode_mesh_qc, "shared_search_mesh_limitation": "all candidate electrodes remain modeled as electrode material; inactive electrodes use a zero-current channel and may alter shunting/local fields; only the winner has a four-active-electrode independent rerun", "parameter_sources": ["raw/fem/simnibs_simulation_20260729-024122.mat", "independent_fem/simnibs_simulation_20260729-030053.mat"], "roi_element_assignment": "gray-matter tetrahedron centroid within 20 mm Euclidean radius of subject-space center (-41, -13, 66) mm"},
        "spatial": {
            "coordinate_system": "ernie individual conform physical space in millimeters",
            "display_convention": "L-R, P-A, and I-S axes in individual conform space; no MNI transform was executed or validated",
            "color_scales": {
                "final_slice_shared": {"vmin_v_per_m": 0.0, "vmax_v_per_m": spatial_extrema["shared_color_scale_v_per_m"][1], "statistic": "joint maximum", "reference_population": "search and final four-active-electrode gray-matter tetrahedra", "reference_data": "figures/spatial_extrema.json", "shared_by_figure_ids": ["four_electrode_field", "optimized_field"], "clipping_rule": "none", "unit": "V/m"},
                "carrier_triptych": {"vmin_v_per_m": 0.0, "vmax_v_per_m": carrier_display_vmax, "statistic": "P99", "reference_population": "strictly positive |E1|, |E2| and TImax values across final gray matter", "reference_data": "independent_fem/independent_timax_recompute.npz", "shared_by_figure_ids": ["carrier_and_timax_field"], "clipping_rule": "display colors capped at P99; machine values unchanged", "unit": "V/m"},
                "surface_field": {"vmin_v_per_m": 0.0, "vmax_v_per_m": surface_display_vmax, "statistic": "P99", "reference_population": "strictly positive final gray-matter surface TImax nodes", "reference_data": "figures/figure_data/fig00_final_surface_context.npz", "shared_by_figure_ids": ["field_surface"], "clipping_rule": "display colors capped at surface P99; machine values unchanged", "unit": "V/m"},
            },
        },
        "protocol": {"current_amplitude_convention": "peak", "current_amplitude_note": "每个载波回路固定输入 +1/-1 mA 峰值；净电流为 0；未优化幅度分配，未验证设备或客户 64 通道帽可执行性", "amplitude_optimization_performed": False, "device_executability_validated": False, "electrodes": [
            {"circuit": "carrier_1", "electrode": "F5", "current_ma": 1.0, "run_id": run_id}, {"circuit": "carrier_1", "electrode": "P5", "current_ma": -1.0, "run_id": run_id},
            {"circuit": "carrier_2", "electrode": "FC3", "current_ma": 1.0, "run_id": run_id}, {"circuit": "carrier_2", "electrode": "CP3", "current_ma": -1.0, "run_id": run_id},
        ]},
        "results": {"primary_field": "TImax maximum envelope amplitude", "value_of_record": "four_active_electrode_fem_recomputation", "search_field_role": "candidate_ranking_only", "value_of_record_rule": "Candidate ranking uses the shared search field. For the selected montage, field values and regional statistics are computed from the four-active-electrode FEM recomputation. Threshold-derived metrics must additionally cite their threshold provenance; the primary fixed threshold is a cross-stage constant derived from the search field.", "field_amplitude_convention": "maximum_envelope_amplitude", "field_amplitude_note": "基场按峰值电流标定；TImax 不是 RMS", "weighted_quantile": {"method_id": "volume_weighted_empirical_left", "definition": "Sort tetrahedra by TImax ascending and select the first value where cumulative tetrahedron volume reaches the requested probability", "weight": "tetrahedron_volume_mm3", "sort_order": "value_ascending", "selection_rule": "first cumulative-volume value reaching probability", "search_side": "left", "interpolation": "none", "boundary_rule": "include the first value at or above the cumulative-volume target", "reported_probabilities": [0.05, 0.5, 0.95]}, "statistical_definitions": {"mean": "tetrahedron-volume weighted", "median_p05_p95": "tetrahedron-volume weighted empirical quantiles; ascending TImax, first cumulative-volume value reaching probability, no interpolation", "max": "unweighted single-tetrahedron maximum", "suprathreshold_volume_mm3": "sum of tetrahedron volumes where TImax >= threshold", "threshold_coverage_pct": "100 * suprathreshold volume / total volume of the same region"}, "threshold": {"value_v_per_m": threshold, "reference_definition": "0.5 * search-stage whole-gray-matter single-tetrahedron maximum", "anchor_statistic": "search-stage whole-gray-matter single-tetrahedron maximum", "anchor_value_v_per_m": 2.0 * threshold, "anchor_is_single_element_extreme": True, "mesh_convergence_assessed": False, "sensitivity_note": "Post-hoc exploratory threshold; sensitivity analysis also uses 0.5 * final four-active-electrode whole-gray-matter maximum", "coverage_unit": "%", "coverage_denominator": "same-region tetrahedron volume", "coverage_definition": "100 * suprathreshold tetrahedron volume / total tetrahedron volume of the same region"}, "threshold_provenance": {"value_v_per_m": threshold, "definition": "0.5 * search-stage whole-gray-matter single-tetrahedron maximum", "timing": "post_hoc_exploratory", "prespecified": False, "limitation": "mesh-extreme dependent; formal studies require a prospectively declared threshold or multi-threshold sensitivity analysis"}, "peak_subject_coordinate_mm": spatial_extrema["four_electrode_recompute_field"]["off_target_peak_conform_mm"], "peak_distance_to_target_mm": spatial_extrema["four_electrode_recompute_field"]["distance_from_roi_center_mm"], "peak_distance_definition": "Euclidean distance between the off-target gray-matter argmax tetrahedron center and ROI center in individual conform space", "peak_inside_target_roi": False, "peak_location_assessment": {"statistic": "off-target gray-matter single-tetrahedron argmax", "is_single_element_argmax": True, "position_stability_assessed": False, "mesh_convergence_assessed": False, "conductivity_sensitivity_assessed": False, "electrode_position_sensitivity_assessed": False, "run_id": validation["run_id"], "interpretation_note": "Potential off-target exposure locator only; not a stable anatomical hotspot"}, "robustness_summary": {"status": "not_assessed", "mesh_convergence_assessed": False, "conductivity_sensitivity_assessed": False, "electrode_position_sensitivity_assessed": False, "interpretation_note": "Threshold sensitivity is reported, but no field re-solve robustness analysis was completed"}, "suprathreshold_gray_matter_volume_mm3": whole_gray["suprathreshold_volume_mm3"], "suprathreshold_gray_matter_fraction_pct": whole_gray["threshold_coverage_pct"], "suprathreshold_gray_matter_denominator_volume_mm3": whole_gray["volume_mm3"], "suprathreshold_gray_matter_fraction_definition": "100 * suprathreshold whole-gray-matter tetrahedron volume / total whole-gray-matter tetrahedron volume", "recompute_half_max_threshold_sensitivity": recompute_threshold_sensitivity, "roi_metrics": roi_metrics, "target_to_offtarget_mean_ratio": validation["independent_target_to_off_target_mean_ratio"], "search_stage_roi_metrics": search_stage_roi_metrics, "search_stage_target_to_offtarget_mean_ratio": winner["target_to_off_target_mean_ratio"]},
        "quality_gates": [
            {"name": "候选集合完整穷举", "hard_gate": True, "status": "pass"},
            {"name": "每回路电流守恒", "hard_gate": True, "status": "pass"},
            {"name": "基场共网格对齐", "hard_gate": True, "status": "pass"},
            {"name": "四活动电极构型复算已完成", "hard_gate": True, "status": "pass", "scope": "computation_completed_only; not mesh-convergence, solver, scientific, or publication validation"},
            {"name": "四活动电极体无共享四面体或三角面且无退化四面体", "hard_gate": True, "status": "pass"},
            {"name": "F5-FC3 两共享边界节点的点接触在正式研究网格中消除", "hard_gate": False, "status": "warning"},
            {"name": "正式研究应事前声明并满足电流校准误差门限", "hard_gate": False, "status": "warning", "observed_range_pct": [0.3, 6.9]},
            {"name": "manifest SHA-256 完整性", "hard_gate": True, "status": "pass"},
        ],
        "legacy_search_stage_result": result,
        "optimization": {
            "objective_function": result["objective"], "constraints": result["constraints"],
            "optimization_mode": "finite_exhaustive_enumeration",
            "search_strategy_provenance": {
                "requested_method": "SimNIBS TesFlexOptimization",
                "native_attempt_status": result["native_tesflex_attempt"]["status"],
                "native_attempt_stage": result["native_tesflex_attempt"]["stage"],
                "native_result_produced": False,
                "fallback_strategy": "manually declared finite candidate library",
                "candidate_bipolar_pairs": ["F5-P5", "F6-P6", "FC3-CP3"],
                "candidate_library_scope": "manual declaration; not an exhaustive enumeration of all 10-10 positions or all 64-channel combinations",
                "global_optimum_claim_allowed": False,
                "interpretation": "Optimal denotes deterministic ranking only within the declared three-pair finite library.",
            },
            "all_admissible_candidates_evaluated": True,
            "candidate_ranking_artifact": "tables/candidate_ranking.csv",
            "search_space": result["search_space"], "solutions": [{**row, "solution_type": "finite_library_candidate"} for row in result["all_candidates"]],
            "selected_solution": winner, "convergence": result["convergence"],
            "four_electrode_recompute_status": "completed",
            "independent_prefix_semantics": "Historical path/key prefix only. In this package, independent_* denotes the selected montage's four-active-electrode field recomputation after removing inactive candidate electrodes; it is not an independent mesh-convergence, solver, or scientific validation.",
            "scientific_validation_status": "formal_study_prerequisites_pending",
            "search_mesh_electrode_qc_status": "not_verified_per_electrode",
            "search_mesh_point_contact_scope": "The shared search mesh uses the same 40 mm F5 and FC3 definitions, but inactive electrodes share a combined zero-current channel tag, so per-electrode topology was not independently audited. Candidate ranking, the search-derived threshold, and Figures 1-6 are exploratory and require re-meshing with positive clearance before publication.",
            "four_electrode_recompute_run_id": validation["run_id"],
            "four_electrode_recompute": {**validation, "status": "completed"},
        },
        "publication_figure_readiness": figure_readiness,
        "references": REFERENCES,
        "figures": [
            {"id": "model_target_registration", "src": "figures/fig00_model_target_registration.png", "alt": "ernie 个体 T1、组织边界与球形灰质 ROI 的三正交定位", "run_id": validation["run_id"]},
            {"id": "stimulation_montage", "src": "figures/fig00_inverse_electrode_montage.png", "alt": "F5-P5 和 FC3-CP3 四电极方案及电流方向", "run_id": validation["run_id"]},
            {"id": "carrier_and_timax_field", "src": "figures/fig00_four_electrode_e1_e2_timax.png", "alt": "最终四电极 E1、E2 与 TImax 的灰质轴向场分布", "run_id": validation["run_id"], "color_scale": {"vmin_v_per_m": 0.0, "vmax_v_per_m": carrier_display_vmax, "statistic": "P99", "reference_population": "strictly positive |E1|, |E2| and TImax across final gray matter", "reference_data": "independent_fem/independent_timax_recompute.npz", "shared_by_figure_ids": ["carrier_and_timax_field"], "clipping_rule": "display only; machine values unchanged", "unit": "V/m"}},
            {"id": "field_surface", "src": "figures/fig00_four_electrode_timax_surface.png", "alt": "最终四电极 TImax 的灰质表面分布", "run_id": validation["run_id"], "color_scale": {"vmin_v_per_m": 0.0, "vmax_v_per_m": surface_display_vmax, "statistic": "P99", "reference_population": "strictly positive final gray-matter surface TImax nodes", "reference_data": "figures/figure_data/fig00_final_surface_context.npz", "shared_by_figure_ids": ["field_surface"], "clipping_rule": "display only; machine values unchanged", "unit": "V/m"}},
            {"id": "four_electrode_field", "src": "figures/fig07_four_electrode_timax_slices.png", "alt": "四活动电极场解复算的 ROI 中心三正交切面和靶外极值切面", "run_id": validation["run_id"]},
            {"id": "four_electrode_quantitative", "src": "figures/fig08_four_electrode_quantitative.png", "alt": "四活动电极场解复算的 P95、极值、覆盖率和超阈体积", "run_id": validation["run_id"]},
            {"id": "optimized_field", "src": "figures/fig01_optimized_timax_slices.png", "alt": "最优离散方案 TImax 的 ROI 中心三正交切面和靶外最大值切面", "run_id": run_id},
            {"id": "candidate_ranking", "src": "figures/fig02_candidate_objective.png", "alt": "三个离散候选的目标区 TImax 均值", "run_id": run_id},
            {"id": "target_offtarget", "src": "figures/fig03_target_offtarget.png", "alt": "目标区和靶外灰质定量比较", "run_id": run_id},
            {"id": "finite_search_order", "src": "figures/fig04_candidate_ranking.png", "alt": "有限候选集合完整排序", "run_id": run_id},
            {"id": "current_constraints", "src": "figures/fig05_current_constraints.png", "alt": "两个载波回路电流约束", "run_id": run_id},
            {"id": "four_electrode_recompute", "src": "figures/fig06_independent_fem.png", "alt": "搜索阶段与四活动电极复算阶段的描述性比较，差异不可归因", "run_id": validation["run_id"]},
        ],
        "artifacts": artifacts,
    }
    report_data["results"].update({
        "value_of_record_status": "demo_only_electrical_independence_not_established",
        "numerical_field_validity": {
            "electrical_independence_established": False,
            "local_shunting_may_affect_e1_e2_timax": True,
            "values_citable_for_formal_research": False,
            "affected_outputs": [
                "E1",
                "E2",
                "TImax",
                "regional_statistics",
                "threshold_coverage",
            ],
            "replacement_requirement": (
                "Verify positive F5-FC3 clearance, rebuild the electrode mesh, rerun both carrier solves, "
                "and replace all field values and derived statistics before formal citation."
            ),
        },
    })
    report_data["results"]["cross_stage_comparison"] = cross_stage_comparison
    (OUT / "report_data.json").write_text(json.dumps(report_data, ensure_ascii=False, indent=2), encoding="utf-8")
    search_peak = spatial_extrema["search_field"]
    recompute_peak = spatial_extrema["four_electrode_recompute_field"]
    shared_color_max = spatial_extrema["shared_color_scale_v_per_m"][1]
    search_peak_coord = ", ".join(f"{value:.2f}" for value in search_peak["off_target_peak_conform_mm"])
    recompute_peak_coord = ", ".join(f"{value:.2f}" for value in recompute_peak["off_target_peak_conform_mm"])
    rows = "".join(f"<tr><td>{c['rank']}</td><td>{html.escape(c['candidate_id'])}</td><td>{c['target_mean_v_per_m']:.4f}</td><td>{c['off_target_mean_v_per_m']:.4f}</td><td>{c['target_to_off_target_mean_ratio']:.3f}</td></tr>" for c in candidates)
    figures = "".join(f'<figure><img src="figures/{name}.png" alt="{alt}"><figcaption>{caption}</figcaption></figure>' for name, alt, caption in (
        ("fig01_optimized_timax_slices", "最优离散方案 TImax 的目标区中心三正交切面和探索性靶外单元极值切面", f"图 1｜共享候选库电极化网格上的最优 TImax 搜索场，不是四活动电极场解复算结果。显示域仅包含灰质四面体，不显示 CSF、颅骨、头皮或其他非灰质组织，因此本图不能用于判断非灰质暴露。前三个面板为 ROI 中心三正交切面，绿色圆表示 20 mm ROI；第四个面板显示靶外灰质单四面体极值 {result['metrics']['off_target']['max']:.4f} V/m，ernie 个体 conform 坐标 ({search_peak_coord}) mm，距 ROI 中心 {search_peak['distance_from_roi_center_mm']:.2f} mm，轴状切面 z={search_peak['off_target_peak_conform_mm'][2]:.2f} mm；白色叉号标记该单元位置。该极值位置未经过网格收敛和位置稳健性分析，只用于提示潜在靶外暴露，不能解释为稳定解剖热点。目标区与靶外灰质 P95 分别为 {result['metrics']['target']['p95']:.4f} 和 {result['metrics']['off_target']['p95']:.4f} V/m。图 1 与图 7 统一使用 0 至 {shared_color_max:.4f} V/m 的线性色标，可直接比较颜色；该上限为两套灰质场的共同最大值，不做截断。源数据：raw/fields/inverse_ti_results.h5、figures/spatial_extrema.json。"),
        ("fig02_candidate_objective", "三个候选方案的搜索阶段目标区 TImax 均值柱状图", "图 2｜搜索阶段候选方案目标函数。所有候选来自真实 FEM 基场；柱高为 ROI 体积加权均值。源数据：tables/candidate_ranking.csv。"),
        ("fig03_target_offtarget", "搜索阶段最优候选的目标区和靶外灰质定量比较", "图 3｜搜索阶段目标区与靶外灰质比较，仅用于候选筛选证据，不得作为最终方案结果图。该图未由四活动电极复算场复核；目标区均值较高，但靶外仍存在更高的单元峰值。源数据：report_data.json / results.search_stage_roi_metrics。"),
        ("fig04_candidate_ranking", "有限候选集合按目标函数排序", "图 4｜完整候选排序。该图不是迭代收敛曲线。源数据：tables/candidate_ranking.csv。"),
        ("fig05_current_constraints", "两个载波回路固定电流输入的守恒核验", "图 5｜固定输入约束核验。每个载波回路固定输入 +1/-1 mA 峰值，净电流为 0；本次未搜索幅度分配，也未验证设备或客户 64 通道帽可执行性。"),
        ("fig06_independent_fem", "共享候选网格搜索场与四活动电极场解复算的目标区和靶外均值比较", f"图 6｜搜索阶段与四活动电极复算阶段的描述性场解对照。搜索网格目标/靶外均值比为 {winner['target_to_off_target_mean_ratio']:.3f}；四电极复算场分别重新计算目标与靶外均值，均值比为 {validation['independent_target_to_off_target_mean_ratio']:.3f}，未跨场拼接指标。复算场 ROI 均值与搜索场相差 {validation['relative_difference_pct']:.2f}%。两套结果使用相同 ROI 单元和权重，只能确认统计域对齐；复算日志记录 0.3%–6.9% 的估计电流校准误差，搜索阶段缺少同口径汇总，因此该差异不能与运行间校准偏差分离，也不能归因于电极构型。本图不是网格收敛或稳健性证据。两个 SimNIBS Msh 的矢量字段名均为 E；independent_timax_recompute.npz 同时交付 E1、E2、TImax、单元中心、体积及目标/靶外掩膜。源数据：independent_fem/independent_validation.json。"),
        ("fig07_four_electrode_timax_slices", "四活动电极场解复算的目标区中心三正交切面和靶外单元极值切面", f"图 7｜四活动电极场解复算的 TImax 空间分布，不是共享候选搜索场。当前 F5-FC3 电极体存在点接触拓扑，本图仅为演示性结果；正式客户研究或论文制图必须先消除点接触并重新网格、求解和替换本图。显示域仅包含灰质四面体；前三个面板为 ROI 中心三正交切面，绿色圆表示 20 mm ROI；第四个面板显示复算场靶外灰质单四面体极值 {validation['independent_metrics']['off_target']['max']:.4f} V/m，ernie 个体 conform 坐标 ({recompute_peak_coord}) mm，距 ROI 中心 {recompute_peak['distance_from_roi_center_mm']:.2f} mm，轴状切面 z={recompute_peak['off_target_peak_conform_mm'][2]:.2f} mm。图 7 与图 1 统一使用 0 至 {shared_color_max:.4f} V/m 的线性色标，可直接比较颜色；该上限为两套灰质场的共同最大值，不做截断。极值位置尚未经过网格收敛和位置稳健性分析，不能解释为稳定解剖热点。源数据：independent_fem/independent_timax_recompute.npz、figures/spatial_extrema.json。"),
        ("fig08_four_electrode_quantitative", "四活动电极场解复算的 P95、单四面体极值、覆盖率和超阈体积", f"图 8｜四活动电极演示性复算的定量图。当前 F5-FC3 电极体存在点接触拓扑；正式客户研究或论文制图必须先消除点接触并重新网格、求解和替换本图及全部数值。左：目标区与靶外灰质 P95、单四面体极值（V/m）；中：固定绝对阈值 {validation['comparison_threshold_v_per_m']:.4f} V/m 下的覆盖率（%）；右：同一阈值下的超阈体积（mm³）。图中场值及区域统计均由四活动电极复算场计算，但阈值常数来自搜索场全灰质单四面体极值的 50%，属于明确保留的跨阶段参数，不是复算场自身 50% 阈值。以复算场自身 50% 极值 {recompute_half_max_threshold:.4f} V/m 进行敏感性重算时，目标区/靶外覆盖率为 {recompute_threshold_sensitivity['target']['threshold_coverage_pct']:.2f}%/{recompute_threshold_sensitivity['off_target']['threshold_coverage_pct']:.2f}%，超阈体积为 {recompute_threshold_sensitivity['target']['suprathreshold_volume_mm3']:.1f}/{recompute_threshold_sensitivity['off_target']['suprathreshold_volume_mm3']:.1f} mm³。两种阈值均为事后探索性定义，且受未完成的网格收敛分析限制。源数据：independent_fem/independent_timax_recompute.npz。"),
    ))
    delivery_rows = "".join(f"<tr><td>{html.escape(item['path'])}</td><td>{item['type']}</td><td>{html.escape(item['purpose'])}</td></tr>" for item in artifacts)
    delivery_rows += "<tr><td>report_data.json</td><td>JSON</td><td>报告结构、ROI 统计、优化约束及图数映射</td></tr><tr><td>manifest.json</td><td>JSON</td><td>全包文件大小与 SHA-256 校验清单</td></tr><tr><td>delivery_root_sha256.txt</td><td>TXT</td><td>manifest.json 的链外根摘要</td></tr>"
    delivery_section = f"<h2>核心交付物索引</h2><p>所有相对路径均以本报告所在目录为起点。HDF5、NIfTI 和 Msh 中的电场单位为 V/m，空间为 ernie 个体 conform 空间；CSV 与 Excel 保留未四舍五入统计值。本表列出核心复算、制图和方法文件；PDF、其余 PNG/SVG、原始 FEM 网格与日志等全部文件及 SHA-256 见 <code>manifest.json</code>。<code>delivery_root_sha256.txt</code> 记录 manifest 的根摘要，不纳入 manifest，以避免循环依赖；正式交付时应通过邮件、项目系统或签收单独立传递该摘要，接收方先核对根摘要，再核对包内文件。</p><div class='scroll'><table><thead><tr><th>文件</th><th>格式</th><th>用途</th></tr></thead><tbody>{delivery_rows}</tbody></table></div>"
    delivery_section += "<p><b>历史文件名前缀。</b><code>independent_fem/</code>、<code>independent_validation.json</code>、<code>independent_timax_recompute.npz</code> 及结构化字段 <code>independent_*</code> 为既有复算链的历史命名。本报告中它们只表示移除未激活候选电极后，对入选方案进行四活动电极场解复算；不表示独立网格收敛验证、独立求解器验证或科学验证通过。正式引用应写作“四活动电极构型复算”。</p><p><b>电极材料模型。</b>搜索与四活动电极复算均使用 SimNIBS 单层简单电极：40 × 40 mm 椭圆、单层厚度 2 mm，材料为 <code>Electrode_rubber</code>，固定电导率 29.4 S/m；未单独建模凝胶、海绵或接触阻抗。参数来自实际 Python 输入、求解日志和 SimNIBS 4.6.0 默认电导率表。客户电极若包含凝胶、海绵、不同材料或接触阻抗，必须按真实层结构重新建模和求解。</p><p><b>搜索场体数据限制。</b><code>ernie_inverse_ti_TImax_discrete_optimum.nii.gz</code> 是共享候选网格的搜索场，只用于候选排序诊断与复核，禁止作为入选方案的最终统计、结果口径或论文结果图。它由 SimNIBS 线性体插值生成，1 mm 各向同性、256 × 256 × 208，覆盖 WM/GM 参考体素的 100%；完整元数据见同名 JSON。</p>"
    delivery_section += "<p><b>四活动电极体数据。</b><code>four_electrode_timax_T1grid.nii.gz</code> 是主要体空间文件：SimNIBS 将最终灰质 TImax 四面体单元值赋到 <code>raw/anatomy/ernie_T1_conform.nii.gz</code> 的 1 mm 网格，两者 shape、体素尺寸和 affine 完全一致，可直接叠加。该文件覆盖 724,014/724,014 个参考灰质体素（100%）；非灰质体素为 0。体素统计受栅格化方式影响，ROI 定量和体积仍以 Msh/NPZ/JSON 的四面体数据为准。<code>four_electrode_timax_conform_nearest_centroid.nii.gz</code> 是质心落栅辅助文件，同体素取最大值，仅覆盖 70.37% 的参考灰质体素，只用于快速定位，不得作为主要体场或定量输入。完整 affine、参考影像哈希、覆盖率和生成方法见 <code>four_electrode_timax_T1grid.json</code>。</p>"
    css = """*{box-sizing:border-box}body{margin:0;font-family:'Microsoft YaHei',sans-serif;color:#17211d;background:#f4f6f5;line-height:1.7}main{width:100%;max-width:1260px;margin:auto;background:white;padding:56px 70px;overflow-wrap:anywhere}h1{font-size:36px}h2{border-top:1px solid #d8dfdb;padding-top:28px;margin-top:42px}img{width:100%;height:auto}figure{margin:28px 0}figcaption{color:#4f5d56;font-size:14px}table{width:100%;border-collapse:collapse;font-size:14px}th,td{border:1px solid #d8dfdb;padding:9px;text-align:left}code{font-size:11px;overflow-wrap:anywhere}.notice{border-left:5px solid #a36b18;background:#fff8e8;padding:18px}.metric{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}.metric div{min-width:0;border:1px solid #d8dfdb;padding:16px}.metric b{display:block;font-size:25px;color:#176b50}@media(max-width:700px){main{padding:24px 18px}h1{font-size:30px;line-height:1.35}.metric{grid-template-columns:1fr}.scroll{max-width:100%;overflow:auto}}@media print{body{background:white}main{max-width:none;padding:20mm}figure{break-inside:avoid}}"""
    document = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>个体化逆向 TI 仿真交付报告</title><style>{css}</style></head><body><main><p>QuanLan BrainScience · 科研演示交付</p><h1>个体化逆向 TI 仿真报告</h1><p>SimNIBS ernie 示例受试者｜受限候选库离散搜索｜报告编号 QLA-SIMNIBS-INVERSE-TI-ERNIE-20260729-01</p><div class="notice"><b>结果范围</b><br>本次结果是三个预定义双极电极对构成的有限候选库内最优，不代表连续头皮位置的全局最优，也未完成客户 64 通道脑电帽的可执行映射。</div><h2>主要结果</h2><div class="metric"><div><span>有限集合最优方案</span><b>{winner['carrier_1']}<br>+ {winner['carrier_2']}</b></div><div><span>搜索阶段 ROI 均值</span><b>{winner['target_mean_v_per_m']:.4f} V/m</b></div><div><span>独立 FEM 复算 ROI 均值</span><b>{validation['independent_rerun_roi_mean_v_per_m']:.4f} V/m</b></div></div><p>在预先声明并完整穷举的三个可行双载波组合中，{winner['carrier_1']} 与 {winner['carrier_2']} 组合获得最高目标区体积加权平均 TImax。搜索阶段的目标/靶外灰质均值比为 {winner['target_to_off_target_mean_ratio']:.3f}；独立 FEM 已在独立网格内重新计算目标区与靶外灰质的均值、P95、均值比和超阈体积。搜索阶段目标区均值高于靶外灰质均值，但靶外灰质最大值为 {result['metrics']['off_target']['max']:.4f} V/m，高于目标区最大值 {result['metrics']['target']['max']:.4f} V/m，因此结果不能表述为全脑热点已限制在目标区内。</p><h2>任务与约束</h2><p>目标区为 ernie 个体空间坐标 (-41, -13, 66) mm 周围 20 mm 球形范围内的灰质单元。目标函数为 ROI 内 TImax 的四面体体积加权均值。每个载波回路使用 +1/-1 mA 峰值电流，净电流为 0；两个回路不得共享电极。</p>{figures}<h2>全部候选（搜索阶段）</h2><div class="scroll"><table><thead><tr><th>排名</th><th>载波组合</th><th>搜索阶段 ROI mean (V/m)</th><th>搜索阶段靶外 mean (V/m)</th><th>搜索阶段均值比</th></tr></thead><tbody>{rows}</tbody></table></div><h2>计算可信度与未完成项</h2><p>三个基场均由 SimNIBS 4.6.0 对 ernie 头模型进行真实有限元求解，并在相同电极化网格上组合。选中方案又在仅包含四个活动电极的独立网格上复算；搜索阶段 ROI 均值为 {validation['search_stage_roi_mean_v_per_m']:.4f} V/m，独立复算为 {validation['independent_rerun_roi_mean_v_per_m']:.4f} V/m，相对差异 {validation['relative_difference_pct']:.2f}%，满足预设 10% 容差。该复算没有验证搜索阶段的靶外统计或均值比。候选集合已完整穷举，因此不需要也不应展示迭代收敛曲线。尚未完成网格收敛、电导率敏感性、电极偏移、客户 MRI 配准和客户 64 通道帽映射；这些项目在正式研究交付前仍属于硬门禁。</p><h2>科学解释边界</h2><p>本报告支持比较所声明候选方案的模型电场分布，不验证神经激活、临床安全性、治疗效果或个体处方。论文使用时应同时报告候选库、目标函数、电流幅度约定、组织电导率、网格信息和未完成的稳健性分析。</p></main></body></html>"""
    document = document.replace(
        "<h2>主要结果</h2>",
        "<h2>主要结果</h2><div class='notice'><b>发布状态</b><br><code>report.status=ready</code> 仅表示演示交付包文件已完整生成，不表示已满足论文发表或客户正式个体研究条件。当前论文发布状态为 <code>formal_study_prerequisites_pending</code>；点接触、客户 MRI 与真实电极几何、非灰质暴露及稳健性分析等项目完成后，必须重新求解并替换结果。</div><div class='notice'><b>搜索阶段的点接触边界</b><br>共享搜索网格沿用直径 40 mm 的 F5 与 FC3 电极定义，但零电流候选共用合并标签，本次未完成搜索网格的逐电极拓扑 QC。四活动电极复算已确认 F5-FC3 存在两节点点接触，因此候选排序、由搜索场极值导出的 0.3245 V/m 主阈值以及图 1–6 均按探索性证据处理，不得作为正式论文结果；正式研究须先验证正的电极间隙并重做搜索、阈值推导和全部图件。</div><div class='notice'><b>引用口径</b><br>共享搜索场用于三个候选的排序，并一次性导出主分析固定阈值。对入选的 F5-P5 + FC3-CP3 方案，mean、P95、单元极值、均值比、空间图及阈值下的区域体积统计均在四活动电极复算场上计算；引用覆盖率和超阈体积时，必须同时注明阈值常数来自搜索场，不能表述为所有参数均由复算场独立产生。本报告另给出复算场自身 50% 极值阈值的敏感性结果。不得跨场拼接分子、分母或场统计量。</div>",
        1,
    )
    document = document.replace("搜索阶段 ROI 均值</span>", "搜索阶段 ROI 均值（仅排序）</span>", 1)
    document = document.replace("独立 FEM 复算 ROI 均值</span>", "四活动电极 ROI 均值（结果口径）</span>", 1)
    search_metric_card = f'<div><span>搜索阶段 ROI 均值（仅排序）</span><b>{winner["target_mean_v_per_m"]:.4f} V/m</b></div>'
    final_metric_card_source = f'<div><span>四活动电极 ROI 均值（结果口径）</span><b>{validation["independent_rerun_roi_mean_v_per_m"]:.4f} V/m</b></div>'
    final_metric_card = f'<div><span>四活动电极 ROI 均值（原始 FEM 估计）</span><b>{validation["independent_rerun_roi_mean_v_per_m"]:.2f} V/m</b><small>未校正；求解日志估计电流校准误差 0.3%–6.9%</small></div>'
    final_ratio_card = f'<div><span>四活动电极目标/靶外均值比</span><b>{validation["independent_target_to_off_target_mean_ratio"]:.3f}</b></div>'
    document = document.replace(search_metric_card + final_metric_card_source, final_metric_card + final_ratio_card, 1)
    document = document.replace(
        final_ratio_card + "</div>",
        final_ratio_card + "</div><div class='notice'><b>演示结果限制</b><br>上述四活动电极数值来自存在 F5-FC3 点接触拓扑的演示网格，不是可直接投稿的正式客户结果。正式研究必须先调整电极尺寸或位置，验证正的电极间隙，并重新网格、求解和替换全部数值及图件。</div>",
        1,
    )
    volume_summary = (
        f"<p><b>搜索阶段（仅用于候选排序）。</b>目标/靶外灰质均值比 {winner['target_to_off_target_mean_ratio']:.3f} 仅表示两个区域的平均 TImax 之比，不是聚焦指数。"
        f"目标 ROI 体积为 {result['metrics']['target']['volume_mm3']:.1f} mm³，靶外灰质体积为 {result['metrics']['off_target']['volume_mm3']:.1f} mm³。"
        f"按固定绝对阈值 {result['metrics']['target']['threshold_v_per_m']:.4f} V/m（由搜索场全灰质单四面体最大值的 50% 一次性导出），目标区超阈体积为 {result['metrics']['target']['suprathreshold_volume_mm3']:.1f} mm³（{result['metrics']['target']['threshold_coverage_pct']:.2f}%），"
        f"靶外灰质超阈体积为 {result['metrics']['off_target']['suprathreshold_volume_mm3']:.1f} mm³（{result['metrics']['off_target']['threshold_coverage_pct']:.2f}%）。"
        f"靶外超阈体积约为目标区的 {result['metrics']['off_target']['suprathreshold_volume_mm3'] / result['metrics']['target']['suprathreshold_volume_mm3']:.2f} 倍；"
        f"全部超阈灰质体积中约 {100 * result['metrics']['target']['suprathreshold_volume_mm3'] / (result['metrics']['target']['suprathreshold_volume_mm3'] + result['metrics']['off_target']['suprathreshold_volume_mm3']):.2f}% 位于目标 ROI。"
        f"<b>四活动电极结果口径：</b>目标区和靶外覆盖率分别为 {validation['independent_metrics']['target']['threshold_coverage_pct']:.2f}% 和 {validation['independent_metrics']['off_target']['threshold_coverage_pct']:.2f}%；靶外/目标超阈体积为 {validation['independent_metrics']['off_target']['suprathreshold_volume_mm3'] / validation['independent_metrics']['target']['suprathreshold_volume_mm3']:.2f} 倍，全部超阈灰质体积中 {100 * validation['independent_metrics']['target']['suprathreshold_volume_mm3'] / (validation['independent_metrics']['target']['suprathreshold_volume_mm3'] + validation['independent_metrics']['off_target']['suprathreshold_volume_mm3']):.2f}% 位于目标 ROI。"
        "两套结果均不能据均值比宣称刺激已聚焦。该阈值由本次结果事后导出且受网格单元极值影响，正式研究需预先声明阈值或报告多阈值敏感性。</p>"
    )
    independent_summary = (
        f"<p><b>四活动电极独立复算。</b>目标区均值 {validation['independent_metrics']['target']['mean']:.4f} V/m，"
        f"靶外灰质均值 {validation['independent_metrics']['off_target']['mean']:.4f} V/m，目标/靶外均值比 {validation['independent_target_to_off_target_mean_ratio']:.3f}；"
        f"目标区与靶外 P95 分别为 {validation['independent_metrics']['target']['p95']:.4f} 和 {validation['independent_metrics']['off_target']['p95']:.4f} V/m；"
        f"目标区单四面体极值为 {validation['independent_metrics']['target']['max']:.4f} V/m，靶外及全灰质单四面体极值为 {validation['independent_metrics']['off_target']['max']:.4f} V/m。"
        f"沿用搜索阶段同一固定绝对阈值 {validation['comparison_threshold_v_per_m']:.4f} V/m；该数值由搜索场一次性导出，相当于独立全灰质单四面体极值的 {100.0 * validation['comparison_threshold_v_per_m'] / validation['independent_metrics']['whole_gray_matter']['max']:.2f}%，不是独立网格极值的 50%。独立网格目标区总体积为 {validation['independent_metrics']['target']['volume_mm3']:.1f} mm³，超阈体积为 {validation['independent_metrics']['target']['suprathreshold_volume_mm3']:.1f} mm³，覆盖率 {validation['independent_metrics']['target']['threshold_coverage_pct']:.2f}%；"
        f"独立网格靶外灰质总体积为 {validation['independent_metrics']['off_target']['volume_mm3']:.1f} mm³，超阈体积为 {validation['independent_metrics']['off_target']['suprathreshold_volume_mm3']:.1f} mm³，覆盖率 {validation['independent_metrics']['off_target']['threshold_coverage_pct']:.2f}%。"
        f"<b>阈值敏感性：</b>若改用四活动电极复算场自身全灰质单四面体极值的 50%（{recompute_half_max_threshold:.4f} V/m），目标区超阈体积为 {recompute_threshold_sensitivity['target']['suprathreshold_volume_mm3']:.1f} mm³、覆盖率 {recompute_threshold_sensitivity['target']['threshold_coverage_pct']:.2f}%；靶外超阈体积为 {recompute_threshold_sensitivity['off_target']['suprathreshold_volume_mm3']:.1f} mm³、覆盖率 {recompute_threshold_sensitivity['off_target']['threshold_coverage_pct']:.2f}%；全部超阈灰质中目标区占 {recompute_threshold_sensitivity['target_share_of_all_suprathreshold_gray_pct']:.2f}%。该敏感性阈值同样是事后定义，只用于说明结论对阈值口径的依赖，不能替代正式研究的预设多阈值分析。"
        f"两套完整电极化网格的单元数不同，但用于统计的 {validation['gray_statistical_domain_alignment']['gray_tetrahedra']:,} 个灰质四面体几何完全对齐：目标/靶外掩膜逐元素一致，灰质单元中心最大差 {validation['gray_statistical_domain_alignment']['gray_centers_max_abs_difference_mm']:.1f} mm，体积最大差 {validation['gray_statistical_domain_alignment']['gray_volumes_max_abs_difference_mm3']:.1f} mm³。因此两套灰质分析域总体积相同是对齐核验结果，不是沿用分母。上述总体积、超阈体积和覆盖率仍均在独立场上成组计算，未与搜索场交叉拼接。"
        f"相对搜索网格，独立网格靶外均值变化 {100.0 * (validation['independent_metrics']['off_target']['mean'] / result['metrics']['off_target']['mean'] - 1.0):+.1f}%，"
        f"目标区超阈体积变化 {100.0 * (validation['independent_metrics']['target']['suprathreshold_volume_mm3'] / result['metrics']['target']['suprathreshold_volume_mm3'] - 1.0):+.1f}%，"
        f"靶外超阈体积变化 {100.0 * (validation['independent_metrics']['off_target']['suprathreshold_volume_mm3'] / result['metrics']['off_target']['suprathreshold_volume_mm3'] - 1.0):+.1f}%；"
        f"靶外/目标超阈体积倍数由 {result['metrics']['off_target']['suprathreshold_volume_mm3'] / result['metrics']['target']['suprathreshold_volume_mm3']:.2f} 变为 {validation['independent_metrics']['off_target']['suprathreshold_volume_mm3'] / validation['independent_metrics']['target']['suprathreshold_volume_mm3']:.2f}。"
        "这些指标不适用 ROI 均值的 10% 容差，不能据 6.25% 的 ROI 均值差异宣称聚焦性或靶外暴露已稳健复现。</p>"
    )
    comparison_specs = [
        ("目标区 mean", result["metrics"]["target"]["mean"], validation["independent_metrics"]["target"]["mean"], "V/m"),
        ("靶外 mean", result["metrics"]["off_target"]["mean"], validation["independent_metrics"]["off_target"]["mean"], "V/m"),
        ("目标区 P95", result["metrics"]["target"]["p95"], validation["independent_metrics"]["target"]["p95"], "V/m"),
        ("靶外 P95", result["metrics"]["off_target"]["p95"], validation["independent_metrics"]["off_target"]["p95"], "V/m"),
        ("目标区单四面体极值", result["metrics"]["target"]["max"], validation["independent_metrics"]["target"]["max"], "V/m"),
        ("靶外单四面体极值", result["metrics"]["off_target"]["max"], validation["independent_metrics"]["off_target"]["max"], "V/m"),
        ("目标区总体积", result["metrics"]["target"]["volume_mm3"], validation["independent_metrics"]["target"]["volume_mm3"], "mm³"),
        ("靶外总体积", result["metrics"]["off_target"]["volume_mm3"], validation["independent_metrics"]["off_target"]["volume_mm3"], "mm³"),
        ("目标区超阈体积", result["metrics"]["target"]["suprathreshold_volume_mm3"], validation["independent_metrics"]["target"]["suprathreshold_volume_mm3"], "mm³"),
        ("靶外超阈体积", result["metrics"]["off_target"]["suprathreshold_volume_mm3"], validation["independent_metrics"]["off_target"]["suprathreshold_volume_mm3"], "mm³"),
        ("目标区覆盖率", result["metrics"]["target"]["threshold_coverage_pct"], validation["independent_metrics"]["target"]["threshold_coverage_pct"], "%"),
        ("靶外覆盖率", result["metrics"]["off_target"]["threshold_coverage_pct"], validation["independent_metrics"]["off_target"]["threshold_coverage_pct"], "%"),
        ("目标/靶外均值比", winner["target_to_off_target_mean_ratio"], validation["independent_target_to_off_target_mean_ratio"], "ratio"),
        ("固定绝对阈值（由搜索场极值 50% 一次性导出）", result["metrics"]["target"]["threshold_v_per_m"], validation["comparison_threshold_v_per_m"], "V/m"),
    ]
    comparison_rows = "".join(
        f"<tr><td>{label}</td><td>{search_value:.4f}</td><td>{independent_value:.4f}</td><td>{unit}</td></tr>"
        for label, search_value, independent_value, unit in comparison_specs
    )
    comparison_rows += (
        f"<tr><td>靶外极值 conform 坐标</td><td>({search_peak_coord})</td><td>({recompute_peak_coord})</td><td>mm</td></tr>"
        f"<tr><td>靶外极值到 ROI 中心距离</td><td>{search_peak['distance_from_roi_center_mm']:.2f}</td>"
        f"<td>{recompute_peak['distance_from_roi_center_mm']:.2f}</td><td>mm</td></tr>"
    )
    comparison_table = (
        "<h2>搜索与独立复算对照</h2><div class='scroll'><table><thead><tr><th>指标</th><th>搜索网格（候选排序）</th>"
        "<th>四活动电极复算场（最优方案结果口径）</th><th>单位</th></tr></thead><tbody>"
        f"{comparison_rows}</tbody></table></div>"
        "<p>表内场强保留四位小数用于复核机器可读结果，不表示具有相同位数的剂量精度。四活动电极求解日志记录 0.3%–6.9% 的估计电流校准误差，原始场值未作事后缩放；该日志量不是统计置信区间，也不能直接换算成场强上下界。正式剂量比较前必须预设并满足校准误差门限。两列分别在各自场解内计算，不得跨列组合分子与分母。10% 仅为预先记录的操作性阈值；6.25% 差异与最高 6.9% 的日志误差同量级，且搜索场缺少同口径误差汇总，因此不能据此认定两场数值一致或稳健复现。均值比仅在所有相关场值受同一比例因子缩放时保持不变；当前未证明各电极误差满足这一条件。</p>"
    )
    cluster_rows = "".join(
        "<tr>"
        f"<td>{cluster['cluster_id']}</td><td>{cluster['volume_mm3']:.1f}</td>"
        f"<td>({', '.join(f'{value:.1f}' for value in cluster['centroid_conform_mm'])})</td>"
        f"<td>{cluster['centroid_distance_from_roi_center_mm']:.1f}</td>"
        f"<td>{cluster['min_tetrahedron_center_distance_from_roi_center_mm']:.1f}–{cluster['max_tetrahedron_center_distance_from_roi_center_mm']:.1f}</td>"
        f"<td>{'是' if cluster['shares_face_with_target_suprathreshold'] else '否'}</td><td>{cluster['max_v_per_m']:.4f}</td>"
        "</tr>"
        for cluster in cluster_analysis["clusters"][:10]
    )
    largest_cluster = cluster_analysis["clusters"][0]
    largest_cluster_sensitivity = cluster_sensitivity["clusters"][0]
    cluster_table = (
        "<h2>四活动电极复算场：靶外超阈空间聚类</h2>"
        f"<p>按四面体共享完整三角面定义连通，固定阈值 {cluster_analysis['threshold_v_per_m']:.4f} V/m 下共得到 "
        f"{cluster_analysis['cluster_count']:,} 个靶外连通域。最大连通域体积 {largest_cluster['volume_mm3']:.1f} mm³，"
        f"占靶外超阈总体积 {100 * largest_cluster['volume_mm3'] / cluster_analysis['total_off_target_suprathreshold_volume_mm3']:.1f}%；"
        f"其体积加权质心为 ({', '.join(f'{value:.1f}' for value in largest_cluster['centroid_conform_mm'])}) mm，"
        f"距 ROI 中心 {largest_cluster['centroid_distance_from_roi_center_mm']:.1f} mm；单元中心距 ROI 中心范围为 "
        f"{largest_cluster['min_tetrahedron_center_distance_from_roi_center_mm']:.1f}–{largest_cluster['max_tetrahedron_center_distance_from_roi_center_mm']:.1f} mm，"
        f"并且{'与' if largest_cluster['shares_face_with_target_suprathreshold'] else '不与'}目标区超阈四面体共享完整三角面。"
        f"因此该主团包含紧邻 20 mm ROI 边界的连续外溢，不能仅按 {largest_cluster['centroid_distance_from_roi_center_mm']:.1f} mm 质心解释为孤立远端热点。下表列体积最大的 10 个连通域；完整 {cluster_analysis['cluster_count']:,} 个连通域及包围盒见 CSV/JSON。</p>"
        f"<p><b>阈值敏感性。</b>上述 {cluster_analysis['threshold_v_per_m']:.4f} V/m 阈值由搜索场极值事后导出。改用四活动电极场自身 50% 极值 {cluster_sensitivity['threshold_v_per_m']:.4f} V/m 后，靶外连通域由 {cluster_analysis['cluster_count']:,} 个变为 {cluster_sensitivity['cluster_count']:,} 个；最大连通域体积为 {largest_cluster_sensitivity['volume_mm3']:.1f} mm³，占该阈值下靶外超阈体积的 {100 * largest_cluster_sensitivity['volume_mm3'] / cluster_sensitivity['total_off_target_suprathreshold_volume_mm3']:.1f}%，单元中心距 ROI 中心范围为 {largest_cluster_sensitivity['min_tetrahedron_center_distance_from_roi_center_mm']:.1f}–{largest_cluster_sensitivity['max_tetrahedron_center_distance_from_roi_center_mm']:.1f} mm，并且{'仍与' if largest_cluster_sensitivity['shares_face_with_target_suprathreshold'] else '不再与'}目标区超阈区共享完整三角面。两种阈值均为事后探索性定义，聚类数、主团体积和空间范围会随阈值改变，不能解释为阈值稳健的解剖外溢边界。完整敏感性结果见带 <code>four_electrode_half_max_sensitivity</code> 后缀的 CSV/JSON。</p>"
        "<div class='scroll'><table><thead><tr><th>聚类</th><th>体积 (mm³)</th><th>质心 conform 坐标 (mm)</th>"
        "<th>质心距 ROI (mm)</th><th>单元中心距 ROI 范围 (mm)</th><th>与目标超阈区共享面</th><th>峰值 (V/m)</th></tr></thead><tbody>"
        f"{cluster_rows}</tbody></table></div>"
        "<p>另交付质心采样二值 NIfTI，用于快速定位和叠加显示；它把每个超阈四面体的质心映射到 T1 体素，不是四面体体积栅格化，不能用该 NIfTI 的体素数重算体积。体积统计以 CSV/JSON 的四面体体积之和为准。</p>"
    )
    document = document.replace("</p><h2>任务与约束</h2>", f"</p>{volume_summary}{comparison_table}{independent_summary}{cluster_table}<h2>任务与约束</h2>", 1)
    document = document.replace(
        "靶外灰质定义为全部灰质减去目标 ROI，不设置 atlas 排除区或过渡缓冲区。",
        "靶外灰质定义为全部灰质减去目标 ROI，不设置 atlas 排除区或过渡缓冲区。区域 mean、median、P05 和 P95 均按四面体体积加权；加权分位数按 TImax 升序累计单元体积，取首次达到 5%、50% 或 95% 区域体积的单元值，不做插值。max 为未加权的单四面体极值；超阈覆盖率为超阈四面体体积之和除以同一区域总体积。",
        1,
    )
    document = document.replace(
        f"搜索阶段目标区均值高于靶外灰质均值，但靶外灰质最大值为 {result['metrics']['off_target']['max']:.4f} V/m，高于目标区最大值 {result['metrics']['target']['max']:.4f} V/m，因此结果不能表述为全脑热点已限制在目标区内。",
        f"搜索阶段目标区均值高于靶外灰质均值，但靶外灰质单个四面体极值为 {result['metrics']['off_target']['max']:.4f} V/m，高于目标区单个四面体极值 {result['metrics']['target']['max']:.4f} V/m；目标区与靶外灰质 P95 分别为 {result['metrics']['target']['p95']:.4f} 和 {result['metrics']['off_target']['p95']:.4f} V/m。单元极值未经过网格收敛或独立靶外复算，不能定位为稳定解剖热点，但足以说明结果不能表述为全脑高场均已限制在目标区内。",
        1,
    )
    document = document.replace(
        "单元极值未经过网格收敛或独立靶外复算，不能定位为稳定解剖热点，",
        "搜索与独立网格均出现靶外单元极值高于目标区极值，但这些极值尚未经过网格收敛，不能定位为稳定解剖热点，",
        1,
    )
    document = document.replace(
        "三个基场均由 SimNIBS 4.6.0 对 ernie 头模型进行真实有限元求解，并在相同电极化网格上组合。",
        "三个基场均由 SimNIBS 4.6.0 对 ernie 头模型进行真实有限元求解，并在相同电极化网格上组合。共享搜索网格保留候选库全部电极几何；未激活电极使用零电流通道，但仍作为电极材料参与模型，可能改变分流和局部场。当前仅对排名第一的组合完成四活动电极独立复算，因此搜索阶段绝对场强及未复算候选的名义排名不能解释为设备构型下已验证的结果。SimNIBS 原生 <code>TesFlexOptimization</code> 曾在头模型准备阶段发生 Windows 原生进程异常，尚未进入优化，也没有产生候选解、目标函数或收敛数值；因此本演示改用预先声明的有限候选库穷举。",
    )
    document = document.replace(
        "该复算没有验证搜索阶段的靶外统计或均值比。",
        "独立 FEM 已在独立网格内重新计算目标区与靶外灰质均值、P95、均值比和超阈体积；这些独立网格指标单独报告，不与搜索网格数值交叉拼接。",
        1,
    )
    document = document.replace(
        "尚未完成网格收敛、电导率敏感性、电极偏移、客户 MRI 配准和客户 64 通道帽映射；这些项目在正式研究交付前仍属于硬门禁。",
        "尚未完成网格收敛、电导率敏感性、电极偏移、CSF/颅骨/头皮等非灰质组织暴露评估、客户 MRI 配准和客户 64 通道帽映射；这些项目在正式研究交付前仍属于硬门禁。",
        1,
    )
    document = document.replace(
        "不验证神经激活、临床安全性、治疗效果或个体处方。",
        "本次定量分析域仅包含灰质，不验证非灰质组织暴露、神经激活、临床安全性、治疗效果或个体处方。",
        1,
    )
    document = document.replace(
        "也未完成客户 64 通道脑电帽的可执行映射。",
        "也未完成客户 64 通道脑电帽的可执行映射。本场解使用直径 40 mm、厚 2 mm 的圆形电极；帽式小触点即使使用同名 10-10 位置，也不能沿用本报告场强、局部电流密度或暴露结论，必须按真实触点尺寸、接触层和位置重新进行 FEM 求解。两个载波回路的电流幅度比固定为 1:1，未搜索幅度分配，因此结果仅表示固定幅度下的候选库内最优；载波频率和差频也未纳入本次电场计算。",
        1,
    )
    document = document.replace(
        "两个回路不得共享电极。",
        "两个回路不得共享电极。TImax 使用 SimNIBS 4.6.0 <code>get_maxTI(E1, E2)</code>：先令 ||E1||≥||E2|| 且夹角 α≤90°；若 ||E2||≤||E1||cosα，则 TImax=2||E2||，否则 TImax=2||E2×(E1−E2)||/||E1−E2||。本报告数值包含因子 2，表示最大调制幅值，不是半幅、RMS 或载波峰峰值。E1、E2 按每回路 1 mA peak 标定；两个回路同时按相同比例缩放时，TImax 在线性准静态假设下按同一比例缩放。设备或文献若采用 RMS、峰峰值或其他电流口径，必须先换算。有限元求解采用准静态、线性场叠加假设；本演示未指定载波频率 f1、f2 或差频 Δf，频率没有参与电场或 TImax 数值计算。正式刺激协议必须另行声明 f1、f2、Δf、相位和设备波形约定。靶外灰质定义为全部灰质减去目标 ROI，不设置 atlas 排除区或过渡缓冲区。均值按四面体体积加权；median、P05 和 P95 为体积加权经验分位数：按 TImax 升序累加四面体体积，取累计体积首次达到目标概率的单元值，不插值。max 是区域内未加权的单四面体极值。固定绝对阈值由搜索场全灰质单四面体极值的 50% 一次性导出，并原样用于两套网格；超阈体积为达到该固定阈值的四面体体积之和，覆盖率分母为同一区域总体积。",
        1,
    )
    document = document.replace(
        "本报告支持比较所声明候选方案的模型电场分布，",
        "本报告支持比较固定 1:1 载波幅度下所声明候选方案的模型电场分布；未优化载波幅度分配，",
        1,
    )
    document = document.replace(
        "尚未完成网格收敛、电导率敏感性、电极偏移、CSF/颅骨/头皮等非灰质组织暴露评估、客户 MRI 配准和客户 64 通道帽映射；",
        "尚未完成网格收敛、电导率敏感性、电极偏移、CSF/颅骨/头皮等非灰质组织暴露评估、ROI 的 atlas 解剖标签与 MNI 坐标映射、个体到标准空间变换交付、客户 MRI 配准，以及按客户 64 通道帽真实触点尺寸、接触层和位置重建的 FEM；位置名称相同不等于帽式方案可执行；",
        1,
    )
    document = document.replace(
        "本次定量分析域仅包含灰质，不验证非灰质组织暴露、神经激活、临床安全性、治疗效果或个体处方。",
        "本次定量分析域仅包含灰质；ROI 仅以 ernie 个体 conform 坐标和球形灰质掩膜定义，未生成 atlas 解剖标签、MNI 坐标或个体到标准空间变换证据，因此不能跨被试复现靶点，也不能直接迁移到客户 MRI。本报告不验证非灰质组织暴露、神经激活、临床安全性、治疗效果或个体处方。",
        1,
    )
    document = document.replace(
        "该极值未经过网格收敛或独立靶外复算，只用于提示潜在靶外暴露，不能解释为稳定解剖热点。",
        "搜索场与独立复算场均出现靶外单四面体极值高于目标区极值；但极值位置尚未经过网格收敛和位置稳健性分析，只用于提示潜在靶外暴露，不能解释为稳定解剖热点。",
        1,
    )
    document = document.replace(
        f"相对差异 {validation['relative_difference_pct']:.2f}%，满足预设 10% 容差。独立 FEM 已在独立网格内重新计算目标区与靶外灰质均值、P95、均值比和超阈体积；",
        f"ROI 均值相对差异 {validation['relative_difference_pct']:.2f}%，数值上低于预先记录的 10% 操作性阈值；但该差异与复算日志最高 6.9% 的估计电流校准误差同量级，且搜索场缺少同口径误差汇总，不能据此认定两场数值一致或稳健复现。独立 FEM 已在独立网格内重新计算目标区与靶外灰质均值、P95、均值比和超阈体积；",
        1,
    )
    document = document.replace(
        f"相对差异 {validation['relative_difference_pct']:.2f}%，满足预设 10% 容差。该复算没有验证搜索阶段的靶外统计或均值比。",
        f"ROI 均值相对差异为 {validation['relative_difference_pct']:.2f}%，数值上低于预先记录的 10% 操作性阈值。四活动电极求解日志的估计电流校准误差最高为 6.9%，与该差异同量级；搜索场又缺少同口径误差汇总。因此，该阈值不构成两场数值一致、可复现或稳健的证据。独立网格内靶外均值为 {validation['independent_metrics']['off_target']['mean']:.4f} V/m、P95 为 {validation['independent_metrics']['off_target']['p95']:.4f} V/m、均值比为 {validation['independent_target_to_off_target_mean_ratio']:.3f}；这些指标同样不能由 10% 操作性阈值推导出稳健性结论。",
        1,
    )
    document = document.replace(
        "</tbody></table></div><h2>计算可信度与未完成项</h2>",
        "</tbody></table></div><p>按预设目标函数的确定性计算值，三个候选依次列为第 1、2、3 名。第 2、3 名的目标函数仅相差约 0.5%，且两者的靶外均值和目标/靶外均值比并不相同；当前未对这两个候选执行独立 FEM 复算，也未估计排序不确定度，因此该名义顺序不能解释为稳健优劣。</p><h2>计算可信度与未完成项</h2>",
        1,
    )
    reproduction_parameters = (
        f"<p><b>复算参数。</b>搜索阶段电极为 SimNIBS 椭圆定义的等轴圆形电极（长短轴均为 40 mm，即直径 40 mm），厚 2 mm、几何面积约 12.566 cm²；每个活动回路 1 mA peak 对应名义平均输入密度 0.0796 mA/cm²。该值仅由输入电流除以几何面积得到，不是局部接触电流密度；本次未执行局部电流密度安全评估。电极化网格含 "
        f"{SEARCH_MESH['nodes']:,} 个节点、{SEARCH_MESH['tetrahedra']:,} 个四面体和 {SEARCH_MESH['triangles']:,} 个三角形。"
        f"四活动电极场解复算沿用同一 ernie 头组织网格，仅重建 F5、P5、FC3、CP3 四个活动电极及其邻近区域；活动电极仍为直径 40 mm 的圆形电极、厚 2 mm，并沿用同一套组织电导率。复算网格含 {INDEPENDENT_MESH['nodes']:,} 个节点、{INDEPENDENT_MESH['tetrahedra']:,} 个四面体和 {INDEPENDENT_MESH['triangles']:,} 个三角形。"
        "SimNIBS 四活动电极求解日志记录的估计电流校准误差为 0.3% 至 6.9%。本演示保留原始场值且未对该误差作事后缩放；搜索场未汇总同口径的估计电流校准误差。复算场最高 6.9% 的估计误差与 6.25% 的 ROI 均值差异同量级，因此 10% 操作性阈值不构成数值一致性或稳健复现证据。正式研究应预先设定可接受门限，并在超限时重新网格或重新求解。"
        "按 SimNIBS 默认 ernie 10-10 坐标，F5-FC3 与 P5-CP3 的欧氏中心距分别为 40.21 mm 和 46.85 mm；减去 40 mm 名义直径后的间隙分别为 0.21 mm 和 6.85 mm。网格级 QC 显示：两对电极均无共享四面体或共享三角面，四个电极体均无退化四面体；P5-CP3 无共享节点、跨电极最小节点距 6.57 mm；F5-FC3 共享 2 个边界节点但不共享面或体，属于点接触拓扑。该点接触不等同于有正间隙的电气独立构型，因此本场解仅作演示，正式研究和设备迁移前必须以更小电极或调整位置重新网格，并验证正的最小间隙。完整证据见 <code>independent_fem/electrode_mesh_qc.json</code> 和 <code>tables/electrode_geometry.csv</code>。"
        "ROI 按四面体重心归属：仅纳入灰质标签且重心到 (-41, -13, 66) mm 的欧氏距离不超过 20 mm 的单元。"
        f"各向同性组织电导率（S/m）：{conductivities_text}。电极为单层 2 mm <code>Electrode_rubber</code>，电导率 29.4 S/m；未建模凝胶、海绵或接触阻抗。原始参数与网格量见 "
        "<code>report_data.json</code>、<code>raw/fem/simnibs_simulation_20260729-024122.mat</code> 和 "
        "<code>raw/fields/inverse_ti_results.h5</code>、<code>independent_fem/independent_validation.json</code> 和两组场解复算 Msh。两份 Msh 的矢量字段名为 <code>E</code>；<code>independent_timax_recompute.npz</code> 另存 E1、E2、TImax、单元中心、体积和 ROI 掩膜，可直接复核 0.2614 V/m。两套结果使用相同 ROI 灰质单元与体积权重，只能确认统计域对齐。复算日志记录 0.3%–6.9% 的估计电流校准误差，搜索阶段缺少同口径汇总，因此 6.25% 不能与运行间校准偏差分离，也不能归因于电极构型；它不是灰质网格离散误差或网格收敛证据。灰质数值误差仍需单独的网格收敛测试量化。</p>"
    )
    document = document.replace("</p><figure>", f"</p>{reproduction_parameters}<figure>", 1)
    document = document.replace("<h2>计算可信度与未完成项</h2>", f"{delivery_section}<h2>计算可信度与未完成项</h2>")
    document = document.replace(
        "选中方案又在仅包含四个活动电极的独立网格上复算；",
        "选中方案又沿用同一 ernie 头组织网格，仅重建四个活动电极及其邻近区域并重新求解场；该步骤是电极构型场解复算，不是网格独立性验证；",
        1,
    )
    document = document.replace("独立 FEM 复算", "四活动电极场解复算")
    document = document.replace("独立 FEM", "四活动电极场解复算")
    document = document.replace("独立网格", "四电极复算网格")
    document = document.replace("搜索与独立复算对照", "搜索场与四活动电极场解复算对照")
    document = document.replace("四活动电极独立复算", "四活动电极场解复算")
    document = document.replace("四活动电极场解复算 已", "四活动电极场解复算已")
    document = document.replace("独立复算为", "四活动电极场解复算为")
    document = document.replace("独立全灰质", "四活动电极复算场全灰质")
    document = document.replace("独立场上", "四活动电极复算场上")
    document = document.replace("独立复算 Msh", "场解复算 Msh")
    document = document.replace("不是独立四活动电极复算场", "不是四活动电极场解复算结果")
    document = document.replace("搜索与四电极复算网格均出现", "搜索场与四活动电极复算场均出现")
    document = document.replace("搜索场与独立复算场均出现", "搜索场与四活动电极复算场均出现")
    document = render_customer_report({
        "result": result,
        "validation": validation,
        "search_peak": search_peak,
        "final_peak": recompute_peak,
        "shared_color_max": shared_color_max,
        "artifacts": artifacts,
        "cluster": cluster_analysis,
        "cluster_sensitivity": cluster_sensitivity,
        "figure_readiness": figure_readiness,
    })
    report_html = OUT / "report.html"
    report_pdf = OUT / "report.pdf"
    report_html.write_text(document, encoding="utf-8")
    edge_candidates = [
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    ]
    edge = next((candidate for candidate in edge_candidates if candidate.is_file()), None)
    if edge is None:
        raise FileNotFoundError("Microsoft Edge is required to regenerate report.pdf from report.html")
    subprocess.run(
        [str(edge), "--headless=new", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={report_pdf}", report_html.as_uri()],
        check=True,
        timeout=120,
    )
    if not report_pdf.is_file() or report_pdf.stat().st_size < 50_000:
        raise RuntimeError("report.pdf was not regenerated correctly")
    previous_size = -1
    stable_checks = 0
    for _ in range(30):
        current_size = report_pdf.stat().st_size
        if current_size == previous_size:
            stable_checks += 1
            if stable_checks >= 4:
                break
        else:
            stable_checks = 0
            previous_size = current_size
        time.sleep(1)
    else:
        raise RuntimeError("report.pdf did not become stable before manifest generation")
    manifest = []
    qa_only = {"manifest.json", "dual_acceptance.json", "delivery_root_sha256.txt"}
    for path in sorted(p for p in OUT.rglob("*") if p.is_file() and p.name not in qa_only):
        manifest.append({"path": path.relative_to(OUT).as_posix(), "bytes": path.stat().st_size, "sha256": sha256(path)})
    (OUT / "manifest.json").write_text(json.dumps({"files": manifest}, ensure_ascii=False, indent=2), encoding="utf-8")
    root_digest = sha256(OUT / "manifest.json")
    (OUT / "delivery_root_sha256.txt").write_text(
        f"manifest.json  SHA256  {root_digest}\n",
        encoding="ascii",
    )


if __name__ == "__main__": main()
