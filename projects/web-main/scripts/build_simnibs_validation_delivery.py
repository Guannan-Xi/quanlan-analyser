"""Build a standalone, auditable SimNIBS solver-validation delivery package."""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import math
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "simnibs_smoke_v2_20260728"
SOLVER_SOURCE = ROOT / "outputs" / "simnibs_smoke_20260728"
SOURCE_MESH = Path(
    r"C:\Users\Administrator\Miniconda3\envs\simnibs_env\Lib\site-packages\simnibs\_internal_resources\testing_files\sphere3.msh"
)
OUTPUT = ROOT / "outputs" / "simnibs_publication_report_v2_20260728"
RUN_ID = "simnibs-sphere-bipolar-1ma-20260728"
REPORT_ID = "QLA-SIMNIBS-VALIDATION-20260728-02"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="315F55")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_sheet(workbook: Workbook, title: str, rows: list[dict], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    style_sheet(sheet)


def build_workbook(
    path: Path,
    overview: list[dict],
    regions: list[dict],
    electrodes: list[dict],
    gates: list[dict],
    artifacts: list[dict],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Overview", overview, ["item", "value", "unit", "source"])
    add_sheet(
        workbook,
        "Region statistics",
        regions,
        [
            "region_tag",
            "field",
            "unit",
            "n",
            "volume_mm3",
            "weighting",
            "mean",
            "median",
            "p95",
            "max",
            "interpretation_boundary",
        ],
    )
    add_sheet(
        workbook,
        "Electrodes",
        electrodes,
        ["electrode", "x_mm", "y_mm", "z_mm", "current_mA", "run_id"],
    )
    add_sheet(
        workbook,
        "Quality gates",
        gates,
        ["gate", "status", "hard_gate", "result", "evidence"],
    )
    add_sheet(
        workbook,
        "Artifact index",
        artifacts,
        ["path", "type", "purpose", "unit_space", "run_id", "sha256"],
    )
    dictionary = [
        {"field": "region_tag", "definition": "SimNIBS mesh material tag; no anatomical label is inferred in this validation run."},
        {"field": "magnE", "definition": "Electric-field vector magnitude, V/m."},
        {"field": "magnJ", "definition": "Current-density vector magnitude, A/m^2."},
        {"field": "volume_mm3", "definition": "Total tetrahedral volume included in the statistic, mm^3."},
        {"field": "weighting", "definition": "tetrahedron_volume: means and quantiles are weighted by tetrahedron volume."},
        {"field": "mean/median/p95", "definition": "Volume-weighted descriptive statistics across finite tetrahedral values in the stated region."},
        {"field": "max", "definition": "Maximum finite tetrahedral value; not volume weighted."},
        {"field": "run_id", "definition": "Identifier linking tables, figures, field mesh and report conclusions to one solver run."},
    ]
    add_sheet(workbook, "Data dictionary", dictionary, ["field", "definition"])
    workbook.save(path)


def artifact(path: str, kind: str, purpose: str, unit_space: str, digest: str) -> dict:
    return {
        "path": path,
        "type": kind,
        "purpose": purpose,
        "unit_space": unit_space,
        "run_id": RUN_ID,
        "sha256": digest,
    }


def table_rows(rows: list[list[str]], numeric: set[int] | None = None) -> str:
    numeric = numeric or set()
    rendered = []
    for row in rows:
        cells = "".join(
            f'<td class="num">{html.escape(str(value))}</td>' if index in numeric else f"<td>{html.escape(str(value))}</td>"
            for index, value in enumerate(row)
        )
        rendered.append(f"<tr>{cells}</tr>")
    return "".join(rendered)


def build_region_svg(path: Path, region_rows: list[dict]) -> None:
    rows = sorted(
        (row for row in region_rows if row["field"] == "magnE"),
        key=lambda row: row["region_tag"],
    )
    metrics = (
        ("median", "Median", "#315F55"),
        ("mean", "Mean", "#6F8F85"),
        ("p95", "P95", "#C09248"),
        ("max", "Maximum", "#9A534D"),
    )
    values = [float(row[key]) for row in rows for key, _, _ in metrics]
    positive = [value for value in values if value > 0]
    if not positive:
        raise ValueError("Region chart requires positive magnE values")

    width, height = 960, 560
    left, right, top, bottom = 92, 28, 58, 96
    plot_width = width - left - right
    plot_height = height - top - bottom
    log_min = math.floor(math.log10(min(positive)))
    log_max = math.ceil(math.log10(max(positive)))
    if log_max == log_min:
        log_max += 1

    def y_position(value: float) -> float:
        fraction = (math.log10(value) - log_min) / (log_max - log_min)
        return top + plot_height * (1 - fraction)

    group_width = plot_width / max(len(rows), 1)
    bar_width = min(30.0, group_width / 5.5)
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img" aria-labelledby="title desc">',
        '<title id="title">Volume-weighted electric-field statistics by mesh tag</title>',
        '<desc id="desc">Grouped logarithmic bar chart of median, mean, P95, and maximum electric-field magnitude for each mesh material tag.</desc>',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<g font-family="Arial, Microsoft YaHei, sans-serif" fill="#24312B">',
    ]
    for exponent in range(log_min, log_max + 1):
        value = 10.0**exponent
        y = y_position(value)
        parts.append(f'<line x1="{left}" y1="{y:.2f}" x2="{width-right}" y2="{y:.2f}" stroke="#D9E0DC" stroke-width="1"/>')
        parts.append(f'<text x="{left-12}" y="{y+5:.2f}" text-anchor="end" font-size="13">10^{exponent}</text>')

    baseline = y_position(10.0**log_min)
    for group_index, row in enumerate(rows):
        center = left + group_width * (group_index + 0.5)
        for metric_index, (key, _, color) in enumerate(metrics):
            value = float(row[key])
            x = center + (metric_index - 1.5) * bar_width - bar_width * 0.42
            y = y_position(value)
            bar_height = max(baseline - y, 1.0)
            parts.append(f'<rect x="{x:.2f}" y="{y:.2f}" width="{bar_width*0.84:.2f}" height="{bar_height:.2f}" fill="{color}"/>')
        parts.append(f'<text x="{center:.2f}" y="{height-bottom+30}" text-anchor="middle" font-size="14">Tag {row["region_tag"]}</text>')

    legend_x = left
    for _, label, color in metrics:
        parts.append(f'<rect x="{legend_x}" y="20" width="14" height="14" fill="{color}"/>')
        parts.append(f'<text x="{legend_x+21}" y="32" font-size="13">{label}</text>')
        legend_x += 142
    parts.extend(
        [
            f'<text x="{width/2:.1f}" y="{height-20}" text-anchor="middle" font-size="14">Mesh material tag</text>',
            f'<text transform="translate(24 {top+plot_height/2:.1f}) rotate(-90)" text-anchor="middle" font-size="14">Electric-field magnitude (V/m; log scale)</text>',
            '</g></svg>',
        ]
    )
    path.write_text("".join(parts), encoding="utf-8")


def write_manifest(source_manifest: dict) -> None:
    manifest_files = sorted(
        path for path in OUTPUT.rglob("*") if path.is_file() and path.name != "manifest.json"
    )
    manifest = {
        "report_id": REPORT_ID,
        "run_id": RUN_ID,
        "artifact_type": "simnibs_solver_validation_delivery",
        "status": "blocked_for_anatomical_or_ti_conclusions",
        "generated_at": datetime.now().astimezone().isoformat(),
        "source_mesh_sha256": source_manifest["mesh"]["sha256"],
        "files": [
            {
                "path": path.relative_to(OUTPUT).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            for path in manifest_files
        ],
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def build_html(stats: dict, source_manifest: dict, region_rows: list[dict], gates: list[dict], artifacts: list[dict]) -> str:
    volume_e = stats["volume_elements"]["magnE"]
    volume_j = stats["volume_elements"]["magnJ"]
    region_e = [row for row in region_rows if row["field"] == "magnE"]
    region_body = table_rows(
        [
            [
                row["region_tag"], row["n"], f'{row["volume_mm3"]:.1f}', f'{row["mean"]:.6f}', f'{row["median"]:.6f}',
                f'{row["p95"]:.6f}', f'{row["max"]:.6f}', row["unit"]
            ]
            for row in region_e
        ],
        {1, 2, 3, 4, 5, 6},
    )
    gate_body = "".join(
        f'<tr><td>{html.escape(row["gate"])}</td><td><span class="status {row["status"]}">{html.escape(row["status_label"])}</span></td><td>{html.escape(row["result"])}</td><td>{html.escape(row["evidence"])}</td></tr>'
        for row in gates
    )
    artifact_body = table_rows(
        [[row["path"], row["type"], row["purpose"], row["unit_space"], row["sha256"]] for row in artifacts]
    )
    generated_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")
    mesh = source_manifest["mesh"]
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width,initial-scale=1" />
  <meta name="report-id" content="{REPORT_ID}" />
  <title>SimNIBS 有限元求解验证报告</title>
  <style>
    :root{{--ink:#17221e;--muted:#5d6a64;--line:#d9e0dc;--paper:#fff;--bg:#edf1ef;--green:#315f55;--green2:#e9f1ee;--amber:#8a5b12;--amber2:#fff5df;--red:#8d3434;--red2:#faeaea;--sans:"Noto Sans CJK SC","Microsoft YaHei",Arial,sans-serif;--serif:"Noto Serif CJK SC","Songti SC",SimSun,serif}}
    *{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;overflow-x:hidden;background:var(--bg);color:var(--ink);font:14px/1.72 var(--sans);letter-spacing:0}}button{{font:inherit}}.toolbar{{position:sticky;top:0;z-index:10;display:flex;align-items:center;gap:14px;padding:10px 22px;border-bottom:1px solid var(--line);background:rgba(255,255,255,.96)}}.toolbar strong{{margin-right:auto}}button{{min-height:38px;padding:7px 13px;border:1px solid var(--green);border-radius:5px;background:var(--green);color:#fff;cursor:pointer}}.paper{{width:min(1180px,calc(100% - 32px));margin:24px auto 60px;background:var(--paper);box-shadow:0 16px 44px rgba(23,34,30,.08)}}.cover{{min-height:490px;padding:58px 68px 46px;border-top:7px solid var(--green);display:flex;flex-direction:column}}.brand{{display:flex;justify-content:space-between;color:var(--muted);font-size:12px}}.kicker{{margin-top:76px;color:var(--green);font-weight:700}}h1{{max-width:760px;margin:10px 0 16px;font:700 37px/1.3 var(--serif);overflow-wrap:anywhere}}.subtitle{{max-width:780px;margin:0;color:var(--muted);font-size:17px}}.cover-grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:24px;margin-top:auto;padding-top:28px;border-top:1px solid var(--line)}}.label{{display:block;color:var(--muted);font-size:12px}}.value{{display:block;margin-top:5px;font-weight:650;overflow-wrap:anywhere}}.banner{{display:flex;align-items:center;gap:12px;padding:14px 68px;border-block:1px solid #ead9b7;background:var(--amber2);color:#674511}}.dot{{width:9px;height:9px;border-radius:50%;background:#b87812}}.layout{{display:grid;grid-template-columns:210px minmax(0,1fr);gap:42px;padding:42px 52px 70px}}.layout article{{min-width:0}}nav{{position:sticky;top:76px;align-self:start;padding-right:20px;border-right:1px solid var(--line)}}nav a{{display:block;padding:7px 0;color:#435049;text-decoration:none}}nav a:hover{{color:var(--green)}}section{{scroll-margin-top:72px;padding-bottom:44px;margin-bottom:42px;border-bottom:1px solid var(--line)}}section:last-child{{border:0}}.no{{color:var(--green);font-size:12px;font-weight:700}}h2{{margin:2px 0 8px;font:700 25px/1.4 var(--serif)}}h3{{margin:0 0 7px;font-size:15px}}.lead{{max-width:820px;margin:0 0 22px;color:var(--muted)}}.callout{{padding:21px 23px;border-left:4px solid var(--green);background:#f4f7f5}}.metric-grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:1px;margin-top:20px;border:1px solid var(--line);background:var(--line)}}.metric{{min-height:110px;padding:15px;background:#fff}}.metric b{{display:block;margin-top:11px;font-size:22px;font-variant-numeric:tabular-nums}}.metric small{{color:var(--muted)}}.facts{{display:grid;grid-template-columns:repeat(2,1fr);margin:0;border-top:1px solid var(--line)}}.fact{{display:grid;grid-template-columns:145px 1fr;gap:10px;padding:10px 0;border-bottom:1px solid var(--line)}}.fact:nth-child(odd){{padding-right:22px}}dt{{color:var(--muted)}}dd{{margin:0;font-weight:550;overflow-wrap:anywhere}}figure{{margin:0 0 28px}}figure img{{display:block;width:100%;height:auto;border:1px solid var(--line);background:#fff}}figcaption{{padding-top:9px;color:#46534d;font-size:12px}}.table-wrap{{overflow:auto;border:1px solid var(--line)}}table{{width:100%;border-collapse:collapse;font-size:12px}}th,td{{padding:9px 10px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}}th{{background:#f1f4f2;color:#405048}}td.num{{text-align:right;font-variant-numeric:tabular-nums}}.status{{display:inline-block;padding:2px 7px;border-radius:999px;font-size:11px;font-weight:700;white-space:nowrap}}.status.pass{{background:var(--green2);color:var(--green)}}.status.warning{{background:var(--amber2);color:var(--amber)}}.status.fail{{background:var(--red2);color:var(--red)}}.note{{margin-top:16px;padding:13px 15px;border:1px solid #e4d4ad;background:#fffaf0;color:#654b1c}}.two-col{{columns:2;column-gap:36px}}.two-col article{{break-inside:avoid;margin-bottom:22px}}.two-col p{{margin:0;color:#46534d}}code{{font-family:Consolas,monospace;font-size:11px}}footer{{display:flex;justify-content:space-between;padding:18px 52px;border-top:1px solid var(--line);color:var(--muted);font-size:11px}}
    @media(max-width:850px){{.cover{{padding:40px 32px}}.cover-grid{{grid-template-columns:repeat(2,1fr)}}.banner{{padding-inline:32px}}.layout{{grid-template-columns:1fr;padding:32px 26px}}nav{{position:static;border:0;border-bottom:1px solid var(--line);padding-bottom:16px;columns:2}}.facts{{grid-template-columns:1fr}}.fact:nth-child(odd){{padding-right:0}}.two-col{{columns:1}}}}
    @media(max-width:560px){{.toolbar{{display:block;padding:11px 12px}}.toolbar strong{{display:block;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;font-size:12px;line-height:1.35}}.toolbar span,.toolbar button{{display:none}}.paper{{width:100%;max-width:100vw;margin:0;box-shadow:none}}.brand span:last-child{{display:none}}h1{{width:calc(100vw - 40px);max-width:calc(100vw - 40px);font-size:26px;line-height:1.35;word-break:break-word;overflow-wrap:anywhere}}.subtitle{{width:calc(100vw - 40px);max-width:calc(100vw - 40px);white-space:normal;line-break:strict;word-break:normal;overflow-wrap:anywhere}}.cover-grid{{width:calc(100vw - 40px)}}.cover-grid,.metric-grid{{grid-template-columns:repeat(2,minmax(0,1fr))}}.cover-grid>div{{min-width:0}}.cover{{width:100%;max-width:100vw;min-width:0;overflow:hidden;padding:34px 20px}}.banner{{padding:12px 20px}}.layout{{padding:26px 18px}}.fact{{grid-template-columns:115px 1fr}}footer{{padding:16px 20px;display:block}}footer span{{display:block;margin-top:4px}}}}
    @media print{{body{{background:#fff}}.toolbar,nav{{display:none}}.paper{{width:100%;margin:0;box-shadow:none}}.layout{{display:block;padding:34px 46px}}section,figure{{break-inside:avoid}}.cover{{min-height:92vh}}}}
  </style>
</head>
<body>
  <div class="toolbar"><strong>QLanalyser · SimNIBS 验证报告</strong><span>{REPORT_ID}</span><button onclick="window.print()">打印 / 导出 PDF</button></div>
  <main class="paper" data-report-id="{REPORT_ID}" data-build="solver-validation-20260728-v2">
    <header class="cover"><div class="brand"><span>QLanalyser</span><span>SIMULATION VALIDATION REPORT</span></div><div class="kicker">正向有限元求解 · 工程验证</div><h1>SimNIBS 有限元求解验证报告</h1><p class="subtitle">本报告记录三层球体模型上的双电极正向求解、场值统计、数值产物和复算索引，用于确认本机 SimNIBS 计算链路可用。</p><div class="cover-grid"><div><span class="label">模型</span><span class="value">合成三层球体</span></div><div><span class="label">刺激</span><span class="value">双电极 +1/-1 mA</span></div><div><span class="label">报告编号</span><span class="value">{REPORT_ID}</span></div><div><span class="label">生成时间</span><span class="value">{generated_at}</span></div></div></header>
    <div class="banner"><span class="dot"></span><strong>求解器验证通过；个体 MRI/TI 结论尚未生成</strong></div>
    <div class="layout"><nav aria-label="报告目录"><a href="#summary">结果摘要</a><a href="#scope">任务范围</a><a href="#model">模型与刺激</a><a href="#field">电场分布</a><a href="#quant">定量结果</a><a href="#qc">质量控制</a><a href="#methods">方法与限制</a><a href="#delivery">交付索引</a></nav><article>
      <section id="summary"><span class="no">01</span><h2>结果摘要</h2><p class="lead">结论仅适用于本次合成球体有限元求解，不外推至个体脑组织、目标脑区刺激效果或临床效应。</p><div class="callout"><h3>SimNIBS 有限元求解、场量写入、读取和可视化链路已完成端到端验证。</h3><p>结果网格包含 E、|E|、J 和 |J| 四个场。按四面体体积加权后，体单元 |E| 的中位数为 {volume_e['median']:.6f} V/m，P95 为 {volume_e['p95']:.6f} V/m。由于没有个体 MRI、解剖 ROI 和第二组载波场，本报告不提供 TI 包络、靶向性、聚焦性或逆向优化结论。</p></div><div class="metric-grid"><div class="metric"><span class="label">体积加权 |E| 均值</span><b>{volume_e['mean']:.6f}</b><small>V/m · {volume_e['volume_mm3']:.0f} mm³</small></div><div class="metric"><span class="label">体积加权 |E| 中位数</span><b>{volume_e['median']:.6f}</b><small>V/m</small></div><div class="metric"><span class="label">体积加权 |E| P95</span><b>{volume_e['p95']:.6f}</b><small>V/m</small></div><div class="metric"><span class="label">体单元 |E| 最大值</span><b>{volume_e['max']:.6f}</b><small>V/m</small></div></div></section>
      <section id="scope"><span class="no">02</span><h2>任务范围与完成状态</h2><p class="lead">本次运行的目标是验证求解器通路，不是形成个体化刺激处方。完成状态按真实输入和产物判断。</p><div class="table-wrap"><table><thead><tr><th>分析内容</th><th>状态</th><th>说明</th></tr></thead><tbody><tr><td>单回路正向有限元求解</td><td><span class="status pass">已完成</span></td><td>合成球体，双电极 +1/-1 mA。</td></tr><tr><td>场值统计与图件导出</td><td><span class="status pass">已完成</span></td><td>输出 E、|E|、J、|J| 及区域统计。</td></tr><tr><td>个体 MRI 头模型</td><td><span class="status fail">未执行</span></td><td>未提供 T1w/T2w MRI 或 m2m_* 头模型。</td></tr><tr><td>解剖 ROI 分析</td><td><span class="status fail">未执行</span></td><td>球体材料标签不等同于脑区标签。</td></tr><tr><td>TI E1/E2 与 TImax</td><td><span class="status fail">未执行</span></td><td>当前只有一个直流回路场解。</td></tr><tr><td>逆向优化与独立复算</td><td><span class="status fail">未执行</span></td><td>未定义目标 ROI、设备约束和优化目标。</td></tr></tbody></table></div></section>
      <section id="model"><span class="no">03</span><h2>模型与刺激方案</h2><p class="lead">输入、坐标和刺激参数均与结果文件保持一一对应。</p><dl class="facts"><div class="fact"><dt>头模型</dt><dd>SimNIBS 内置三层球体测试网格</dd></div><div class="fact"><dt>网格规模</dt><dd>{mesh['nodes']} 节点，{mesh['elements']} 单元，{mesh['volume_elements']} 体单元</dd></div><div class="fact"><dt>软件</dt><dd>SimNIBS {source_manifest['software']['simnibs']}；Python {source_manifest['software']['python']}</dd></div><div class="fact"><dt>求解环境</dt><dd>{html.escape(source_manifest['software']['platform'])}</dd></div><div class="fact"><dt>电极中心</dt><dd>(95, 0, 0) mm；(-95, 0, 0) mm</dd></div><div class="fact"><dt>施加电流</dt><dd>+1.0 mA；-1.0 mA</dd></div><div class="fact"><dt>电极模型</dt><dd>20 mm 椭圆电极</dd></div><div class="fact"><dt>run_id</dt><dd><code>{RUN_ID}</code></dd></div></dl></section>
      <section id="field"><span class="no">04</span><h2>电场空间分布</h2><p class="lead">SimNIBS 将四面体单元场插值到三个正交平面。上排显示全模型，下排显示内层 tag 3；每排三幅切面共用色标，便于比较空间分布。</p><figure><img src="figures/solver_smoke_test_magnE.png" alt="合成球体全模型和内层区域的轴位、冠状位及矢状位电场强度分布" /><figcaption><b>图 1｜合成球体模型的电场强度分布。</b> A-C 为全模型 z=0、y=0 和 x=0 mm 切面，轮廓线表示材料界面，圆点标记电极中心；D-F 仅显示内层 tag 3。两个色标分别以对应显示域的 P99 为上限，饱和值不代表计算最大值。单位为 V/m。图件以 300 dpi PNG 提供。</figcaption></figure><figure><img src="figures/weighted_region_metrics.svg" alt="不同材料区域标签的体积加权电场描述性统计" /><figcaption><b>图 2｜不同材料区域标签的体积加权 |E| 统计。</b> 柱形依次表示中位数、均值、P95 和最大值，前三项按四面体体积加权，最大值为有限体单元峰值；纵轴采用对数尺度。材料标签不代表解剖脑区。</figcaption></figure></section>
      <section id="quant"><span class="no">05</span><h2>定量结果</h2><p class="lead">均值、中位数和 P95 按四面体体积加权。区域表用于核验材料分区和数值范围，不作为目标 ROI 或靶外脑区比较。</p><div class="table-wrap"><table><thead><tr><th>区域标签</th><th>体单元数</th><th>体积 (mm³)</th><th>加权均值</th><th>加权中位数</th><th>加权 P95</th><th>最大值</th><th>单位</th></tr></thead><tbody>{region_body}</tbody></table></div><div class="note">体积加权 |J|：均值 {volume_j['mean']:.6f} A/m²，中位数 {volume_j['median']:.6f} A/m²，P95 {volume_j['p95']:.6f} A/m²；有限体单元最大值 {volume_j['max']:.6f} A/m²。完整精度见 Excel 和 CSV。</div></section>
      <section id="qc"><span class="no">06</span><h2>数值质量与发布门控</h2><p class="lead">“求解成功”和“可形成正式个体化结论”是两个独立门控。前者已通过，后者因关键输入缺失而阻断。</p><div class="table-wrap"><table><thead><tr><th>门控</th><th>状态</th><th>结果</th><th>证据</th></tr></thead><tbody>{gate_body}</tbody></table></div></section>
      <section id="methods"><span class="no">07</span><h2>方法、解释边界与限制</h2><div class="two-col"><article><h3>有限元求解</h3><p>采用 SimNIBS 4.6.0 对三层球体网格执行双电极正向 tDCS 求解。电极中心位于 x 轴两端，施加电流为 +1 mA 与 -1 mA。结果网格保存矢量电场 E、标量 |E|、矢量电流密度 J 和标量 |J|。</p></article><article><h3>体积加权统计</h3><p>根据 SimNIBS 网格计算每个四面体的体积。均值、中位数和第 95 百分位数以四面体体积为权重；最大值为有限体单元峰值。统计同时按材料标签分层，没有执行假设检验或组间统计。</p></article><article><h3>图像显示</h3><p>SimNIBS 将不连续的单元场以线性方式插值到 321 × 321 规则平面网格，并保持组织边界两侧的场不连续。每排三幅正交切面共用线性色标；色标上限为对应显示域的 P99。区域统计图采用对数纵轴。</p></article><article><h3>限制</h3><p>本次没有个体 MRI、组织分割、解剖 ROI、电极帽配准、第二载波场、TI 包络、稳健性扰动或逆向优化。因此无法回答刺激是否覆盖特定脑区、是否聚焦、是否可形成设备方案或是否具有临床意义。</p></article><article><h3>复核与再计算</h3><p>源球体网格、结果网格、参数表、统计脚本、交付生成脚本与 SHA-256 随包交付，可用于结果复核和二次分析。原始求解日志及当次求解脚本未保留，因此本包不声称能够逐字节复现本次 FEM 求解；重新计算时应生成新的 report_id 和 run_id。</p></article><article><h3>结论边界</h3><p>本报告仅支持“当前计算环境可完成指定球体模型的有限元求解与结果导出”。它不支持神经调控效应、疗效、安全性或个体治疗方案判断。</p></article></div></section>
      <section id="delivery"><span class="no">08</span><h2>数据交付与二次分析索引</h2><p class="lead">Excel 用于快速筛选与制图；CSV 保留机器可读表格；Msh 保存完整场；JSON 与 manifest 提供字段和版本追踪。</p><div class="table-wrap"><table><thead><tr><th>文件</th><th>类型</th><th>用途</th><th>单位 / 空间</th><th>SHA-256</th></tr></thead><tbody>{artifact_body}</tbody></table></div></section>
    </article></div><footer><span>科研工程验证结果，不作为临床诊断、疗效或个体刺激方案依据。</span><span>{REPORT_ID}</span></footer>
  </main>
</body>
</html>"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--finalize-only", action="store_true")
    args = parser.parse_args()
    stats = json.loads((SOURCE / "solver_smoke_test_statistics.json").read_text(encoding="utf-8"))
    source_manifest = json.loads((SOURCE / "manifest.json").read_text(encoding="utf-8"))
    if args.finalize_only:
        if not OUTPUT.exists():
            raise FileNotFoundError(f"Delivery directory does not exist: {OUTPUT}")
        write_manifest(source_manifest)
        print(OUTPUT)
        return
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    for directory in ("figures", "tables", "fields", "methods", "methods/scripts"):
        (OUTPUT / directory).mkdir(parents=True, exist_ok=True)

    copy_map = {
        SOLVER_SOURCE / "solver_smoke_test_magnE.png": OUTPUT / "figures" / "solver_smoke_test_magnE.png",
        SOLVER_SOURCE / "sphere_bipolar_1mA_scalar.msh": OUTPUT / "fields" / "sphere_bipolar_1mA_scalar.msh",
        SOLVER_SOURCE / "sphere_bipolar_1mA_el_currents.geo": OUTPUT / "fields" / "sphere_bipolar_1mA_el_currents.geo",
        SOURCE_MESH: OUTPUT / "fields" / "source_sphere3.msh",
        ROOT / "scripts" / "simnibs_smoke_report.py": OUTPUT / "methods" / "scripts" / "simnibs_smoke_report.py",
        ROOT / "scripts" / "build_simnibs_validation_delivery.py": OUTPUT / "methods" / "scripts" / "build_simnibs_validation_delivery.py",
    }
    for source, target in copy_map.items():
        shutil.copy2(source, target)

    region_rows = []
    for tag, fields in stats["volume_by_tissue_tag"].items():
        for field_name, values in fields.items():
            region_rows.append({
                "region_tag": int(tag),
                "field": field_name,
                "unit": stats["units"][field_name],
                **values,
                "interpretation_boundary": "Mesh material region; not an anatomical ROI.",
            })
    region_rows.sort(key=lambda row: (row["region_tag"], row["field"]))
    build_region_svg(OUTPUT / "figures" / "weighted_region_metrics.svg", region_rows)
    write_csv(
        OUTPUT / "tables" / "region_field_statistics.csv",
        region_rows,
        ["region_tag", "field", "unit", "n", "volume_mm3", "weighting", "mean", "median", "p95", "max", "interpretation_boundary"],
    )
    electrodes = [
        {"electrode": "E1", "x_mm": 95.0, "y_mm": 0.0, "z_mm": 0.0, "current_mA": 1.0, "run_id": RUN_ID},
        {"electrode": "E2", "x_mm": -95.0, "y_mm": 0.0, "z_mm": 0.0, "current_mA": -1.0, "run_id": RUN_ID},
    ]
    write_csv(OUTPUT / "tables" / "electrode_currents.csv", electrodes, list(electrodes[0]))

    gates = [
        {"gate": "有限元求解", "status": "pass", "status_label": "通过", "hard_gate": True, "result": "完成装配、求解和结果网格写入", "evidence": "fields/sphere_bipolar_1mA_scalar.msh"},
        {"gate": "结果完整性", "status": "pass", "status_label": "通过", "hard_gate": True, "result": "E、magnE、J、magnJ 均存在", "evidence": "manifest.json / report_data.json"},
        {"gate": "原始求解日志", "status": "warning", "status_label": "待补充", "hard_gate": False, "result": "运行成功，但本次交付未保留完整控制台日志", "evidence": "结果网格及文件时间戳"},
        {"gate": "个体解剖输入", "status": "fail", "status_label": "阻断", "hard_gate": True, "result": "无个体 MRI 或 m2m_* 头模型", "evidence": "输入清单"},
        {"gate": "TI 场完整性", "status": "fail", "status_label": "阻断", "hard_gate": True, "result": "缺少独立 E1、E2 和 TImax", "evidence": "场清单"},
        {"gate": "正式研究发布", "status": "fail", "status_label": "阻断", "hard_gate": True, "result": "不得用于个体化、临床或论文主结果", "evidence": "本报告限制章节"},
    ]
    write_csv(
        OUTPUT / "tables" / "quality_gates.csv",
        gates,
        ["gate", "status", "status_label", "hard_gate", "result", "evidence"],
    )

    methods = """# 方法与限制\n\n本次运行采用 SimNIBS 4.6.0，在合成三层球体网格上执行 +1/-1 mA 双电极正向 tDCS 有限元求解。\n\n均值、中位数和第 95 百分位数按四面体体积加权。切面图由 SimNIBS 将四面体单元场插值到规则平面网格。由于未使用个体 MRI、解剖 ROI 和第二组载波场，本次结果不得解释为个体化 TI 仿真、靶向性评估或临床结论。\n\n源网格、结果网格及派生统计脚本已交付；原始求解日志和当次求解脚本未保留，因此本包支持结果复核与二次分析，但不声称逐字节复现本次 FEM 求解。\n"""
    (OUTPUT / "methods" / "methods_and_limitations.md").write_text(methods, encoding="utf-8")
    health = f"""# 数据健康检查\n\n- 结果网格：{source_manifest['mesh']['nodes']} 个节点，{source_manifest['mesh']['elements']} 个单元。\n- 体单元：{source_manifest['mesh']['volume_elements']}。\n- 场：{', '.join(source_manifest['mesh']['fields'])}。\n- magnE/magnJ 有限值计数均与单元统计一致，体积加权描述性统计已导出。\n- 区域标签来自网格，不推断解剖名称。\n- 严重限制：没有个体解剖、ROI、TI 双载波和稳健性数据，因此正式研究门控为阻断。\n"""
    (OUTPUT / "methods" / "00_data_health.md").write_text(health, encoding="utf-8")

    preliminary_artifacts = []
    artifact_specs = [
        ("figures/solver_smoke_test_magnE.png", "PNG", "三切面电场强度图", "mm; V/m"),
        ("figures/weighted_region_metrics.svg", "SVG", "体积加权区域描述性统计图", "V/m"),
        ("tables/region_field_statistics.csv", "CSV", "区域场值统计与二次制图", "V/m; A/m^2"),
        ("tables/electrode_currents.csv", "CSV", "电极坐标与电流", "mm; mA"),
        ("tables/quality_gates.csv", "CSV", "质量门控记录", "status"),
        ("fields/sphere_bipolar_1mA_scalar.msh", "Gmsh MSH", "完整有限元场结果", "model mm; SI fields"),
        ("fields/sphere_bipolar_1mA_el_currents.geo", "Gmsh GEO", "电极电流可视化辅助文件", "model mm"),
        ("fields/source_sphere3.msh", "Gmsh MSH", "SimNIBS 原始三层球体网格", "model mm"),
        ("methods/00_data_health.md", "Markdown", "数据健康与限制检查", "n/a"),
        ("methods/methods_and_limitations.md", "Markdown", "方法和解释边界", "n/a"),
        ("methods/scripts/simnibs_smoke_report.py", "Python", "场统计、图件和基础 manifest 生成", "n/a"),
        ("methods/scripts/build_simnibs_validation_delivery.py", "Python", "HTML、Excel 和交付清单生成", "n/a"),
    ]
    for relative, kind, purpose, unit_space in artifact_specs:
        preliminary_artifacts.append(artifact(relative, kind, purpose, unit_space, sha256(OUTPUT / relative)))

    overview = [
        {"item": "report_id", "value": REPORT_ID, "unit": "", "source": "report"},
        {"item": "run_id", "value": RUN_ID, "unit": "", "source": "solver run"},
        {"item": "model", "value": "synthetic three-layer sphere", "unit": "", "source": "SimNIBS test mesh"},
        {"item": "magnE_volume_weighted_mean", "value": stats["volume_elements"]["magnE"]["mean"], "unit": "V/m", "source": "result mesh; tetrahedron-volume weighted"},
        {"item": "magnE_volume_weighted_median", "value": stats["volume_elements"]["magnE"]["median"], "unit": "V/m", "source": "result mesh; tetrahedron-volume weighted"},
        {"item": "magnE_volume_weighted_p95", "value": stats["volume_elements"]["magnE"]["p95"], "unit": "V/m", "source": "result mesh; tetrahedron-volume weighted"},
        {"item": "magnE_volume_max", "value": stats["volume_elements"]["magnE"]["max"], "unit": "V/m", "source": "result mesh"},
    ]
    workbook_path = OUTPUT / "simnibs_validation_results.xlsx"
    build_workbook(workbook_path, overview, region_rows, electrodes, gates, preliminary_artifacts)
    preliminary_artifacts.append(artifact(workbook_path.name, "XLSX", "二次统计、筛选和制图", "mixed; see data dictionary", sha256(workbook_path)))

    report_data = {
        "schema_version": "simnibs.report.v2",
        "analysis_mode": "forward_single",
        "stimulation_modality": "tes",
        "report": {
            "report_id": REPORT_ID,
            "title": "SimNIBS 有限元求解验证报告",
            "status": "blocked",
            "status_label": "求解器验证通过；正式研究报告阻断",
            "organization": "QLanalyser",
            "scientific_question": "本机 SimNIBS 是否能够完成有限元求解并输出可追溯场数据？",
            "generated_at": datetime.now().astimezone().isoformat(),
            "build_identity": "solver-validation-20260728-v2",
        },
        "subject": {"id": "synthetic-sphere3"},
        "target": {"name": "未定义", "definition": "求解器验证不设置解剖 ROI"},
        "model": {
            "mri_source": "未使用",
            "head_model": "SimNIBS 三层球体测试网格",
            "mesh": f"{source_manifest['mesh']['nodes']} nodes; {source_manifest['mesh']['elements']} elements",
            "conductivity_profile": "由本次 SimNIBS 球体求解配置定义；不映射个体组织",
        },
        "spatial": {"coordinate_system": "球体模型笛卡尔坐标，mm", "display_convention": "x/y/z 正交切面"},
        "software": {"simnibs_version": source_manifest["software"]["simnibs"], "python_version": source_manifest["software"]["python"]},
        "protocol": {"electrode_model": "20 mm ellipse", "electrodes": [
            {"circuit": "single", "electrode": "E1", "current_ma": 1.0, "frequency_hz": None, "phase_deg": None, "run_id": RUN_ID},
            {"circuit": "single", "electrode": "E2", "current_ma": -1.0, "frequency_hz": None, "phase_deg": None, "run_id": RUN_ID},
        ]},
        "summary": {"heading": "正向有限元求解与结果导出已完成", "conclusion": "仅支持求解器通路验证；不支持个体化、TI、ROI 或临床结论。"},
        "results": {"primary_field": "magnE", "headline_metrics": overview[3:], "roi_metrics": []},
        "robustness": [],
        "quality_gates": [{"name": row["gate"], "status": row["status"], "hard_gate": row["hard_gate"], "check": row["result"], "result": row["result"], "evidence": row["evidence"]} for row in gates],
        "methods": [
            {"title": "有限元求解", "text": "SimNIBS 4.6.0，三层球体，双电极 +1/-1 mA 正向 tDCS 求解。"},
            {"title": "统计", "text": "均值、中位数和 P95 按四面体体积加权；最大值为有限体单元峰值。"},
            {"title": "限制", "text": "未使用个体 MRI、解剖 ROI、TI 双载波、稳健性分析或逆向优化；原始求解日志及当次求解脚本未保留。"},
        ],
        "figures": [
            {"id": "field_primary", "src": "figures/solver_smoke_test_magnE.png", "alt": "球体模型电场三切面", "run_id": RUN_ID},
            {"id": "roi_distribution", "src": "figures/weighted_region_metrics.svg", "alt": "体积加权材料区域电场统计", "run_id": RUN_ID},
        ],
        "artifacts": preliminary_artifacts,
    }
    report_data_path = OUTPUT / "report_data.json"
    report_data_path.write_text(json.dumps(report_data, ensure_ascii=False, indent=2), encoding="utf-8")
    preliminary_artifacts.append(artifact(report_data_path.name, "JSON", "报告结构化数据", "mixed; SI where stated", sha256(report_data_path)))

    report_html = build_html(stats, source_manifest, region_rows, gates, preliminary_artifacts)
    report_path = OUTPUT / "report.html"
    report_path.write_text(report_html, encoding="utf-8")

    write_manifest(source_manifest)
    print(OUTPUT)


if __name__ == "__main__":
    main()
